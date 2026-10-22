"""The shared file readers. Every scan and the quote checker import these.

One function per format, each returning the text at a locator, so a scan can
quote it and the quote checker can find it again. No scan writes its own
reader (ROADMAP.md, "readers").

Locators, as FRAMEWORKS.md writes them:

    PDF          "page 2", "slide 3", "clause 7" (clause: the whole document)
    xlsx         "ARR_Build!E14" or a defined name such as "ARR_TOTAL"
    CSV          "row 118" (row 1 is the header row)
    Markdown     "line 12"
    video        "seg 4", or "t=35-41" (every segment that overlaps the span)

Paths are relative to the room folder unless absolute.
"""

from __future__ import annotations

import csv
import json
import re
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber

from schema.claim import Source

ROOT = Path(__file__).resolve().parent.parent
CACHE_DIR = ROOT / "cache"


class LocatorError(ValueError):
    """The locator does not fit the file, or points past its end."""


# PDF ----------------------------------------------------------------------


@lru_cache(maxsize=64)
def _pdf_pages(path: str) -> tuple[str, ...]:
    with pdfplumber.open(path) as pdf:
        return tuple((page.extract_text() or "") for page in pdf.pages)


def pdf_pages(path: str | Path) -> list[str]:
    """Every page's text, page 1 first."""
    return list(_pdf_pages(str(path)))


def pdf_page(path: str | Path, page: int) -> str:
    """The text of one page, counting from 1. A deck's slide n is page n."""
    pages = _pdf_pages(str(path))
    if not 1 <= page <= len(pages):
        raise LocatorError(f"{path} has {len(pages)} pages, not page {page}")
    return pages[page - 1]


# xlsx ---------------------------------------------------------------------


@lru_cache(maxsize=16)
def _workbook(path: str):
    # data_only reads the values Excel cached, not the formulas.
    return openpyxl.load_workbook(path, data_only=True)


def xlsx_cell(path: str | Path, ref: str) -> Any:
    """The cached value at "Sheet!E14", or at a defined name such as "ARR_TOTAL"."""
    wb = _workbook(str(path))
    if "!" in ref:
        sheet, cell = ref.split("!", 1)
        sheet = sheet.strip("'")
        if sheet not in wb.sheetnames:
            raise LocatorError(f"{path} has no sheet {sheet!r}")
        return wb[sheet][cell.replace("$", "")].value
    if ref not in wb.defined_names:
        raise LocatorError(f"{path} has no defined name {ref!r}")
    sheet, cell = next(iter(wb.defined_names[ref].destinations))
    return wb[sheet][cell.replace("$", "")].value


def xlsx_defined_name(path: str | Path, name: str) -> str:
    """Where a defined name points, as "Sheet!E14"."""
    wb = _workbook(str(path))
    if name not in wb.defined_names:
        raise LocatorError(f"{path} has no defined name {name!r}")
    sheet, cell = next(iter(wb.defined_names[name].destinations))
    return f"{sheet}!{cell.replace('$', '')}"


def xlsx_sheet(path: str | Path, sheet: str) -> list[list[Any]]:
    """Every row of one sheet as cached values, row 1 first."""
    wb = _workbook(str(path))
    if sheet not in wb.sheetnames:
        raise LocatorError(f"{path} has no sheet {sheet!r}")
    return [list(row) for row in wb[sheet].iter_rows(values_only=True)]


def xlsx_sheet_names(path: str | Path) -> list[str]:
    return list(_workbook(str(path)).sheetnames)


# CSV ----------------------------------------------------------------------


def csv_rows(path: str | Path) -> list[dict[str, str]]:
    """The data rows as dicts. The first dict is file row 2."""
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def csv_row_text(path: str | Path, row: int) -> str:
    """One file row as written, counting from 1; row 1 is the header."""
    with open(path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    if not 1 <= row <= len(lines):
        raise LocatorError(f"{path} has {len(lines)} rows, not row {row}")
    return lines[row - 1]


# Markdown -----------------------------------------------------------------


def md_line(path: str | Path, line: int) -> str:
    """One line, counting from 1."""
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    if not 1 <= line <= len(lines):
        raise LocatorError(f"{path} has {len(lines)} lines, not line {line}")
    return lines[line - 1]


# Transcript ---------------------------------------------------------------


def transcript_path(video: str | Path) -> Path:
    """The transcript sits beside the video: pitch_video.mp4 -> pitch_video.transcript.json."""
    video = Path(video)
    return video.with_name(video.stem + ".transcript.json")


def transcript_segments(path: str | Path) -> list[dict]:
    """Segments as {"id", "start_s", "end_s", "text"}.

    Takes our own shape ({"segments": [...]}) or whisper.cpp's -oj output
    ({"transcription": [{"offsets": {"from": ms, "to": ms}, "text": ...}]}).
    A video path is read through its transcript beside it.
    """
    path = Path(path)
    if path.suffix.lower() in {".mp4", ".mov", ".m4a", ".wav"}:
        path = transcript_path(path)
    data = json.loads(path.read_text(encoding="utf-8"))
    if "segments" in data:
        return [
            {"id": int(s.get("id", i)), "start_s": float(s["start_s"]), "end_s": float(s["end_s"]), "text": s["text"].strip()}
            for i, s in enumerate(data["segments"])
        ]
    return [
        {"id": i, "start_s": s["offsets"]["from"] / 1000, "end_s": s["offsets"]["to"] / 1000, "text": s["text"].strip()}
        for i, s in enumerate(data["transcription"])
    ]


def transcript_segment(path: str | Path, seg_id: int) -> dict:
    for seg in transcript_segments(path):
        if seg["id"] == seg_id:
            return seg
    raise LocatorError(f"{path} has no segment {seg_id}")


# Saved copies of outside sources -------------------------------------------


def cached_copy(uri: str, cache_dir: Path = CACHE_DIR) -> Path | None:
    """The saved copy of a URL or API call, through cache/index.json."""
    index = cache_dir / "index.json"
    if not index.exists():
        return None
    entries = json.loads(index.read_text(encoding="utf-8"))
    rel = entries.get(uri)
    if isinstance(rel, dict):
        rel = rel.get("path")
    return (cache_dir / rel) if rel else None


# One entry point ----------------------------------------------------------

_NUM = re.compile(r"(\d+)")


def _locator_number(locator: str) -> int:
    m = _NUM.search(locator)
    if not m:
        raise LocatorError(f"locator {locator!r} has no number")
    return int(m.group(1))


def text_at(source: Source, room_dir: str | Path) -> str:
    """The text a source's locator points at, as a string the quote checker searches."""
    if source.kind in ("url", "api"):
        copy = cached_copy(source.uri)
        if copy is None or not copy.exists():
            raise LocatorError(f"no saved copy of {source.uri} in cache/")
        return copy.read_text(encoding="utf-8", errors="replace")

    path = Path(source.uri)
    if not path.is_absolute():
        path = Path(room_dir) / path
    if not path.exists():
        raise LocatorError(f"{path} does not exist")
    loc = source.locator.strip()
    suffix = path.suffix.lower()

    if source.kind == "video" or suffix in {".mp4", ".mov", ".m4a", ".wav"}:
        segs = transcript_segments(path)
        if loc.startswith("seg"):
            return transcript_segment(path, _locator_number(loc))["text"]
        m = re.match(r"t=(\d+(?:\.\d+)?)(?:-(\d+(?:\.\d+)?))?", loc)
        if not m:
            raise LocatorError(f"video locator {loc!r} is neither 'seg N' nor 't=A-B'")
        start = float(m.group(1))
        end = float(m.group(2) or m.group(1))
        return " ".join(s["text"] for s in segs if s["end_s"] >= start and s["start_s"] <= end)

    if suffix == ".pdf":
        if loc.startswith("clause"):
            return "\n".join(pdf_pages(path))
        return pdf_page(path, _locator_number(loc))

    if suffix in {".xlsx", ".xlsm"}:
        value = xlsx_cell(path, loc)
        return "" if value is None else str(value)

    if suffix == ".csv":
        return csv_row_text(path, _locator_number(loc))

    if suffix in {".md", ".txt"}:
        return md_line(path, _locator_number(loc))

    if suffix == ".json":
        return path.read_text(encoding="utf-8")

    raise LocatorError(f"no reader for {suffix} files")
