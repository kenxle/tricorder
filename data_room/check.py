#!/usr/bin/env python3
"""The room's test. Reads the generated files back and checks every planted flag and decoy.

Run after generate.py, with the room's virtual environment:

    data_room/.venv/bin/python data_room/check.py
    data_room/.venv/bin/python data_room/check.py --room data_room/hollis_voice_clean --key data_room/answer_key_clean.json
    data_room/.venv/bin/python data_room/check.py --room data_room/hollis_voice_variant --key <path to the sealed variant key>

It reads an answer key and the files of the room it names (the main room and answer_key.json by default).
For each flag the key lists, it checks that the shown value is where the key says it is, then recomputes the
true value from the raw files (the way a diligence scan would) and checks it matches the key. For each flag
the key does not list, it checks that the room is honest on that point. It also checks that the cash
reconciles and that the room does not look built to be caught. It exits with a non-zero status on any miss.

Workbooks are read with openpyxl. Formula cells are read for their stored values, which xlsxwriter writes.
PDFs are read with pdfplumber. CSV files are read with the csv module.
"""
import argparse
import csv
import json
import re
import statistics
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import pdfplumber

HERE = Path(__file__).resolve().parent
_ap = argparse.ArgumentParser(description="Check a generated Hollis Voice data room against its answer key.")
_ap.add_argument("--room", default=str(HERE / "hollis_voice"), help="the room folder (default: data_room/hollis_voice)")
_ap.add_argument("--key", default=None, help="the answer key (default: answer_key.json, or none if --room is not the main room)")
_args = _ap.parse_args()
ROOM = Path(_args.room).expanduser().resolve()
if _args.key is None:
    if ROOM != (HERE / "hollis_voice").resolve():
        sys.exit("Give --key with --room.")
    _args.key = str(HERE / "answer_key.json")
KEY = json.loads(Path(_args.key).expanduser().read_text(encoding="utf-8"))
SEED_DROP = ROOM.parent / f"{ROOM.name}_seed_drop"
PLANTS = set(KEY["plants"])
ALL_FLAGS = ["PF1", "PF2", "PF2b", "PF3", "PF4", "PF5", "PF6", "PF7"]
FACTS = KEY["facts"]

PASSES, FAILS = [], []


def ok(name, cond, detail=""):
    (PASSES if cond else FAILS).append(name)
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  [{detail}]" if detail and not cond else ""))
    return cond


def near(a, b, tol_abs=0.0, tol_rel=0.0):
    return abs(a - b) <= max(tol_abs, tol_rel * abs(b))


# ---- readers -----------------------------------------------------------------------------------
_wb_cache, _pdf_cache = {}, {}


def workbook(rel, data_only=True):
    k = (rel, data_only)
    if k not in _wb_cache:
        _wb_cache[k] = openpyxl.load_workbook(ROOM / rel, data_only=data_only)
    return _wb_cache[k]


def named(wb, name):
    d = wb.defined_names[name]
    sheet, cell = d.attr_text.split("!")
    return wb[sheet.strip("'")][cell.replace("$", "")].value


def pdf_pages(rel):
    if rel not in _pdf_cache:
        with pdfplumber.open(ROOM / rel) as pdf:
            _pdf_cache[rel] = [(p.extract_text() or "") for p in pdf.pages]
    return _pdf_cache[rel]


def squash(t):
    return re.sub(r"\s+", " ", t)


def read_csv(rel):
    with open(ROOM / rel, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def num(s):
    return float(str(s).replace("$", "").replace(",", ""))


# ---- files exist, page counts, footer, no em dashes --------------------------------------------
def check_files():
    pages = {"00_welcome/cover_note_and_index.pdf": 1, "01_overview/pitch_deck.pdf": 14, "01_overview/company_overview.pdf": 1,
             "04_customers/order_form_cedar_pine.pdf": 3, "04_customers/pilot_agreement_marlow.pdf": 3,
             "04_customers/order_form_brookline_row.pdf": 2, "05_team/founder_bios.pdf": 2, "06_product/product_one_pager.pdf": 1}
    for rel in KEY["full_room_files"]:
        ok(f"file exists: {rel}", (ROOM / rel).is_file())
    for rel, n in pages.items():
        got = len(pdf_pages(rel))
        ok(f"page count {rel} is {n}", got == n, f"got {got}")
        ok(f"footer on every page of {rel}", all("Synthetic data room for Deal Flow NYC 2026" in p for p in pdf_pages(rel)))
    dash = []
    for rel in KEY["full_room_files"]:
        if rel.endswith(".pdf"):
            text = "\n".join(pdf_pages(rel))
        elif rel.endswith(".xlsx"):
            text = "\n".join(str(c.value) for ws in workbook(rel).worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str))
        else:
            text = (ROOM / rel).read_text(encoding="utf-8")
        if "—" in text:
            dash.append(rel)
    ok("no em dashes in any room file", not dash, str(dash))
    for rel in ("02_financials/financial_model.xlsx", "02_financials/cap_table.xlsx"):
        wb = workbook(rel)
        ok(f"footer on every sheet of {rel}", all(any(isinstance(c.value, str) and c.value.startswith("Synthetic data room") for row in ws.iter_rows() for c in row)
                                                  for ws in wb.worksheets))
    for rel in KEY["seed_drop_files"]:
        a, b = ROOM / rel, SEED_DROP / rel
        ok(f"seed drop copy identical: {rel}", b.is_file() and a.read_bytes() == b.read_bytes())
    ok("seed drop holds only the first-tier files",
       sorted(str(p.relative_to(SEED_DROP)) for p in SEED_DROP.rglob("*") if p.is_file()) == sorted(KEY["seed_drop_files"]))


def check_formulas_have_values():
    for rel in ("02_financials/financial_model.xlsx", "02_financials/cap_table.xlsx"):
        wf, wv = workbook(rel, False), workbook(rel, True)
        n, empty = 0, []
        for ws in wf.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        n += 1
                        if wv[ws.title][c.coordinate].value is None:
                            empty.append(f"{ws.title}!{c.coordinate}")
        ok(f"{rel}: {n} formula cells all carry stored values", n > 0 and not empty, str(empty[:5]))


# ---- the key's own locations --------------------------------------------------------------------
def check_locations():
    for entry in KEY["flags"] + KEY["decoys"]:
        for loc in entry["locations"]:
            tag = f"{entry['id']} location: {loc['file']} {loc.get('sheet', '')}{loc.get('cell', '')}{loc.get('text', '')[:40] if 'text' in loc else ''}"
            rel = loc["file"]
            if rel.endswith(".xlsx") and "cell" in loc:
                wb = workbook(rel)
                cell = loc["cell"]
                if ":" in cell:
                    vals = [c.value for row in wb[loc["sheet"]][cell] for c in row]
                    ok(tag, all(isinstance(v, (int, float)) for v in vals) and len(vals) == 12)
                    continue
                val = wb[loc["sheet"]][cell].value
                if "named_cell" in loc:
                    ok(tag + " (named cell points here)", named(wb, loc["named_cell"]) == val)
                exp = loc.get("expect")
                if exp is None:
                    ok(tag, val is not None)
                elif isinstance(exp, str):
                    ok(tag, val == exp, f"got {val!r}")
                else:
                    ok(tag, isinstance(val, (int, float)) and near(val, exp, 1e-6, 1e-6), f"got {val!r} want {exp!r}")
            elif rel.endswith(".pdf") and "page" in loc:
                text = squash(pdf_pages(rel)[loc["page"] - 1])
                ok(tag, squash(loc["text"]) in text, f"page {loc['page']} lacks the text")
            elif rel.endswith(".md"):
                text = squash((ROOM / rel).read_text(encoding="utf-8"))
                ok(tag, squash(loc["text"]) in text)
            elif rel.endswith(".csv") and "customer_name" in loc:
                rows = [r for r in read_csv(rel) if r["customer_name"] == loc["customer_name"] and r["month"] == "2026-09"]
                ok(tag, len(rows) == 1 and rows[0]["product_type"] == loc["product_type"] and int(rows[0]["mrr"]) == loc["mrr"]
                   and rows[0]["customer_id"] == loc["customer_id"])


# ---- raw-file readers used by the flag checks ---------------------------------------------------
def cust_rows():
    return read_csv("03_metrics/customer_mrr_by_month.csv")


def claims():
    out = {}
    for line in pdf_pages("01_overview/pitch_deck.pdf")[13].splitlines():
        m = re.match(r"^(.*?)\s+(\$?[\d,\.]+%?(?: months| \(12 months\))?|(?:under )?\d+% of ARR)\s+(\d+)$", line.strip())
        if m:
            out[m.group(1)] = (m.group(2), int(m.group(3)))
    return out


def flag(fid):
    return next(f for f in KEY["flags"] if f["id"] == fid)


def month_index(label):
    y, m = label.split("-")
    return (int(y) - 2025) * 12 + int(m) - 1


def check_claims_page():
    got = claims()
    want = {c["label"]: (c["value"], c["slide"]) for c in KEY["deck_claims"]}
    ok("claims page lists every claim in the key", got == want, f"missing {set(want) - set(got)}")
    # each claim's slide holds the number (where the slide shows it in the same form)
    deck = pdf_pages("01_overview/pitch_deck.pdf")
    for label, (val, slide) in want.items():
        if label in ("ARR", "Calls handled, September 2026", "Largest customer", "Net revenue retention", "Clinics on Hollis"):
            if label == "ARR":
                probe = f"${num(val) / 1e6:.3g}M"
            elif label == "Largest customer":
                probe = val.replace("under ", "")
            else:
                probe = val.split(" ")[0]
            ok(f"claims row '{label}' appears on slide {slide}", probe in deck[slide - 1], probe)


def check_pf1():
    f = flag("PF1")
    wb = workbook("02_financials/financial_model.xlsx")
    arr_total = named(wb, "ARR_TOTAL")
    rows = cust_rows()
    sub = sum(int(r["mrr"]) for r in rows if r["month"] == "2026-09" and r["product_type"] == "Subscription")
    pilot = sum(int(r["mrr"]) for r in rows if r["month"] == "2026-09" and r["product_type"] == "Pilot")
    ok("PF1 true ARR from Subscription rows matches the key", 12 * sub == f["true"]["value"], f"{12 * sub}")
    ok("PF1 model ARR_TOTAL is the shown value", arr_total == f["shown"]["value"])
    ok("PF1 gap between ARR_TOTAL and Subscription ARR is the pilot at 12 months", arr_total - 12 * sub == 12 * pilot == f["shown"]["details"]["pilot_arr_counted"])
    ok("PF1 model ARR exceeds true ARR by more than the tolerance", (arr_total - 12 * sub) / (12 * sub) > f["tolerance"])
    m = {r["month"]: r for r in read_csv("03_metrics/monthly_metrics.csv")}
    ok("PF1 monthly_metrics arr (2026-09) counts the pilot", int(m["2026-09"]["arr"]) == arr_total)
    t = squash("\n".join(pdf_pages("04_customers/pilot_agreement_marlow.pdf")))
    ok("PF1 pilot agreement: six-month term and $90,000 in clause 2", "2. Term and fees." in t and "six (6) months" in t and "$90,000" in t)
    ok("PF1 pilot agreement: 30 days notice in clause 7", "7. Termination." in t and "thirty (30) days written notice" in t)
    ok("PF1 pilot agreement: no duty to convert in clause 8", "8. No duty to convert." in t)
    ok("PF1 pilot fees contracted equal the key", 15000 * 6 == f["true"]["details"]["pilot_fees_contracted_total"])


def check_pf2():
    f = flag("PF2")
    wb = workbook("02_financials/financial_model.xlsx")
    arr_total = named(wb, "ARR_TOTAL")
    deck = num(claims()["ARR"][0])
    ok("PF2 deck claims-page ARR is the shown value", deck == f["shown"]["value"])
    ok("PF2 model ARR_TOTAL is the true value (the model figure)", arr_total == f["true"]["value"])
    gap = deck - arr_total
    ok("PF2 gap to the model is as keyed", gap == f["shown"]["details"]["gap_to_model"], str(gap))
    ok("PF2 gap exceeds the 1 percent tolerance", gap / arr_total > f["tolerance"])


def check_pf2b():
    f = flag("PF2b")
    text = (ROOM / "01_overview/pitch_video_script.md").read_text(encoding="utf-8")
    m = re.search(r"\*\*0:35[^\n]*\n\n(.*?)\n\n\*\*", text, re.S)
    ok("PF2b script has a 0:35 section", bool(m))
    para = m.group(1) if m else ""
    phrase = FACTS["arr_video_phrase"]
    spoken = f["shown"]["value"] if f"about {phrase} dollars in annual recurring revenue" in para else None
    ok(f"PF2b spoken ARR is about {phrase} dollars", spoken == f["shown"]["value"])
    deck = num(claims()["ARR"][0])
    ok("PF2b spoken ARR differs from the deck beyond the tolerance", spoken is not None and abs(spoken - deck) / deck > f["tolerance"] and deck == f["true"]["value"])
    words = len(" ".join(l for l in text.splitlines() if l and not l.startswith(("**", "#", "Read this", "---", "Synthetic"))).split())
    ok("PF2b script is about 300 words", 250 <= words <= 340, str(words))


def cohorts_from_csv():
    by = defaultdict(dict)
    for r in cust_rows():
        if r["product_type"] == "Subscription":
            by[r["customer_id"]][month_index(r["month"])] = int(r["mrr"])
    res = []
    for s in range(0, 9):
        group = [c for c, v in by.items() if min(v) == s]
        start = sum(by[c][s] for c in group)
        end = sum(by[c].get(s + 12, 0) for c in group)
        capped = sum(min(by[c].get(s + 12, 0), by[c][s]) for c in group)
        early, kept = [], []
        for c in group:
            last = max(by[c])
            churn = last + 1 if last < 20 else None
            (early if churn is not None and churn - s <= 2 else kept).append(c)
        res.append({"n": len(group), "early": len(early), "start": start, "end": end, "capped": capped,
                    "k_n": len(kept), "k_start": sum(by[c][s] for c in kept), "k_end": sum(by[c].get(s + 12, 0) for c in kept)})
    return res


def check_pf3():
    f = flag("PF3")
    wb = workbook("02_financials/financial_model.xlsx")
    res = cohorts_from_csv()
    ts = sum(r["start"] for r in res)
    nrr, grr = sum(r["end"] for r in res) / ts, sum(r["capped"] for r in res) / ts
    n_true = sum(r["n"] for r in res)
    ok("PF3 rebuilt net retention matches the key", near(nrr, f["true"]["value"], f["tolerance"] / 5), f"{nrr:.4f}")
    ok("PF3 rebuilt gross retention matches the key", near(grr, f["true"]["details"]["gross_revenue_retention"], f["tolerance"] / 5), f"{grr:.4f}")
    ok("PF3 rebuilt customer count matches the key", n_true == f["true"]["details"]["customers"])
    ok("PF3 Cohorts tab customer total is the shown count", named(wb, "COHORT_CUSTOMERS") == f["shown"]["details"]["customers"])
    ok("PF3 Cohorts tab net retention is the shown value", near(named(wb, "NRR_12M"), f["shown"]["value"], f["tolerance"] / 5))
    tab = wb["Cohorts"]
    first_col = [tab.cell(row=r, column=1).value for r in range(5, 14)]
    ok("PF3 nine cohorts are listed", len(first_col) == 9 and first_col[0] == "2025-01" and first_col[-1] == "2025-09")
    ok("PF3 tab drops the early leavers cohort by cohort", all(tab.cell(row=5 + i, column=2).value == r["k_n"] for i, r in enumerate(res)))
    ok("PF3 dropped customers equal the early leavers", n_true - sum(r["k_n"] for r in res) == f["true"]["details"]["customers_dropped"] == sum(r["early"] for r in res))
    ok("PF3 deck slide 6 shows the shown retention", f"{f['shown']['value'] * 100:.1f}%" in pdf_pages("01_overview/pitch_deck.pdf")[5])
    ok("PF3 gap between shown and true net retention is material", f["shown"]["value"] - f["true"]["value"] > 0.1)


def hires_from_csv():
    out = []
    for r in read_csv("05_team/hiring_plan.csv"):
        y, m, _ = r["start_date"].split("-")
        t = (int(y) - 2026) * 12 + int(m) - 9
        out.append((t, float(r["base_salary_usd"])))
    return out


def rerun_cash(wb, hires):
    a = lambda n: named(wb, n)
    load = a("LOAD_FACTOR")
    payroll = a("STAFF_NOW") * a("AVG_BASE_NOW") * load / 12
    rev, cash, burns, runway = a("REV_START_USD_MONTH"), a("CASH_START"), [], None
    for t in range(1, 61):
        payroll += sum(b for (st, b) in hires if st == t) * load / 12
        rev *= 1 + a("REV_GROWTH")
        cogs = rev * (a("TELEPHONY_PCT") + a("LICENCES_PCT"))
        prog = a("PROGRAMS_USD_MONTH") if t >= a("SERIES_A_PLAN_MONTH") else 0
        burn = payroll + a("HOSTING_USD_MONTH") + a("RENT_ETC_USD_MONTH") + prog + cogs - rev
        if t == a("SERIES_A_PLAN_MONTH"):
            cash += a("SERIES_A_RAISE")
        cash -= burn
        burns.append(burn)
        if cash <= 0 and runway is None:
            runway = t - 1
    return (60 if runway is None else runway), sum(burns[:12]) / 12


def check_pf4():
    f = flag("PF4")
    wb = workbook("02_financials/financial_model.xlsx")
    hires = hires_from_csv()
    ok("PF4 hiring_plan.csv lists 24 hires", len(hires) == f["true"]["details"]["hires_in_plan"] == 24)
    ok("PF4 model Hiring_Plan counts 8 hires", named(wb, "HIRES_COUNTED") == f["shown"]["details"]["hires_counted"] == 8)
    ws = wb["Hiring_Plan"]
    model_roles = [ws.cell(row=r, column=1).value for r in range(5, 13)]
    csv_roles = [r["role"] for r in read_csv("05_team/hiring_plan.csv")][:8]
    ok("PF4 the model's eight hires are the first eight in the plan", model_roles == csv_roles)
    runway_true, burn_true = rerun_cash(wb, hires)
    runway_model, burn_model = rerun_cash(wb, hires[:8])
    ok("PF4 runway rerun with all 24 hires matches the key", runway_true == f["true"]["value"], str(runway_true))
    ok("PF4 average burn rerun with all 24 hires matches the key", near(burn_true, f["true"]["details"]["avg_net_burn_first_12_months"], 1.0), f"{burn_true:.0f}")
    ok("PF4 rerun with the model's 8 hires reproduces Cash_Runway", runway_model == named(wb, "RUNWAY_MONTHS") == f["shown"]["value"], str(runway_model))
    ok("PF4 model average burn matches the key", near(named(wb, "AVG_NET_BURN_12M"), f["shown"]["details"]["avg_net_burn_first_12_months"], 1.0))
    claim = claims()
    ok("PF4 deck claims-page runway is the shown value", claim["Runway"][0] == f"{f['shown']['value']} months")
    ok("PF4 deck claims-page burn is the shown value", num(claim["Average monthly net burn, next 12 months"][0]) == f["shown"]["details"]["avg_net_burn_first_12_months"])
    ok("PF4 deck says 24 planned hires", claim["Planned hires, next 12 months"][0] == "24")
    ok("PF4 runway gap is material", f["shown"]["value"] - f["true"]["value"] >= 12)


def check_pf5():
    f = flag("PF5")
    rows = [r for r in cust_rows() if r["month"] == "2026-09" and r["product_type"] == "Subscription"]
    mrr = sorted((int(r["mrr"]) for r in rows), reverse=True)
    total = sum(mrr)
    top1, top3 = mrr[0] / total, sum(mrr[:3]) / total
    ok("PF5 top customer share matches the key", near(top1, f["true"]["value"], 0.0005), f"{top1:.4f}")
    ok("PF5 top-three share matches the key", near(top3, f["true"]["details"]["top3_share"], 0.0005), f"{top3:.4f}")
    big_id = f["true"]["details"]["largest_customer_id"]
    big = [r for r in rows if r["customer_id"] == big_id]
    ok("PF5 the largest customer is Cedar & Pine, with the id the key gives", len(big) == 1 and int(big[0]["mrr"]) == mrr[0]
       and big[0]["locations"] == str(FACTS["big_customer_locations"]) and big[0]["customer_name"] == "Cedar & Pine Animal Hospitals")
    ok("PF5 deck claims the largest customer is under 9%", claims()["Largest customer"][0] == "under 9% of ARR" and top1 > f["shown"]["value"])


def check_pf6():
    f = flag("PF6")
    wb = workbook("02_financials/financial_model.xlsx")
    led = defaultdict(float)
    for r in read_csv("02_financials/expense_ledger.csv"):
        if r["month"] == "2026-09":
            led[r["category"]] += float(r["amount_usd"])
    basis = int(next(r for r in read_csv("03_metrics/monthly_metrics.csv") if r["month"] == "2026-09")["mrr"])
    shown = 1 - (led["telephony"] + led["third_party_licences"]) / basis
    true = 1 - (led["telephony"] + led["third_party_licences"] + led["cloud_hosting"] + led["model_inference"] + led["support_staff"]) / basis
    ok("PF6 ledger gross margin with only telephony and licences is the shown 80%", near(shown, f["shown"]["value"], 0.0005), f"{shown:.4f}")
    ok("PF6 ledger gross margin with all delivery costs matches the key", near(true, f["true"]["value"], f["tolerance"] / 2), f"{true:.4f}")
    ok("PF6 model GM_PCT is the shown value", near(named(wb, "GM_PCT"), f["shown"]["value"], 0.0005))
    ws = wb["P&L"]
    labels = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 40)}
    ok("PF6 hosting and inference sit under R&D", labels.get("Cloud hosting and model inference") == "R&D")
    ok("PF6 support specialists sit under S&M", labels.get("Customer support specialists") == "S&M")
    ok("PF6 deck slide 5 shows 80% gross margin", "80%" in pdf_pages("01_overview/pitch_deck.pdf")[4] and claims()["Gross margin"][0] == "80%")
    ok("PF6 gap between shown and true is material", f["shown"]["value"] - f["true"]["value"] > 0.15)
    # P&L ties to the ledger for the last actual month
    col = 2 + 20 + 1
    tel_row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "Telephony")
    ok("P&L telephony for 2026-09 ties to the ledger", near(ws.cell(row=tel_row, column=col).value, led["telephony"], 0.005))


def check_pf7():
    f = flag("PF7")
    rows = read_csv("03_metrics/monthly_metrics.csv")
    cum = [int(r["accounts_created_cumulative"]) for r in rows]
    paying = [int(r["paying_accounts"]) for r in rows]
    ok("PF7 accounts_created_cumulative never decreases", all(b >= a for a, b in zip(cum, cum[1:])))
    ok("PF7 the deck number is not the paying count", num(claims()["Clinics on Hollis"][0]) != paying[-1])
    ok("PF7 the cumulative column's name says cumulative", "cumulative" in "accounts_created_cumulative" and cum[-1] > paying[-1])
    ok("PF7 deck clinic count equals the last cumulative value", num(claims()["Clinics on Hollis"][0]) == cum[-1] == f["shown"]["value"])
    ok("PF7 paying accounts now match the key", paying[-1] == f["true"]["value"])
    ok("PF7 the two differ", cum[-1] != paying[-1])
    sub = {r["customer_id"] for r in cust_rows() if r["month"] == "2026-09" and r["product_type"] == "Subscription"}
    ok("PF7 customer file agrees on paying accounts", len(sub) == paying[-1])


def check_decoys():
    wb = workbook("02_financials/financial_model.xlsx")
    ws = wb["ARR_Build"]
    row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == "Onboarding fees")
    ok("DC1 onboarding fees row adds nothing to ARR", ws.cell(row=row, column=5).value == 0)
    ok("DC1 the row label does not say the fee is excluded", "exclud" not in str(ws.cell(row=row, column=1).value).lower())
    fees = [r for r in cust_rows() if r["product_type"] == "Onboarding"]
    ok("DC1 onboarding rows carry fees and no MRR", fees and all(int(r["mrr"]) == 0 and int(r["one_time_fee"]) == 500 for r in fees))
    sub_arr = 12 * sum(int(r["mrr"]) for r in cust_rows() if r["month"] == "2026-09" and r["product_type"] == "Subscription")
    ok("DC1 the Subscription ARR rebuild is unaffected by onboarding rows", sub_arr == FACTS["arr_true_recurring"])

    dc2 = next(d for d in KEY["decoys"] if d["id"] == "DC2")
    loc = next(l for l in dc2["locations"] if "customer_name" in l)
    months = [r for r in cust_rows() if r["customer_name"] == loc["customer_name"] and r["product_type"] == "Subscription"]
    idx = sorted(month_index(r["month"]) for r in months)
    ok("DC2 Brookline Row has an unbroken run of Subscription months at one price, through 2026-09",
       idx and idx[-1] == 20 and idx == list(range(idx[0], 21)) and len({r["mrr"] for r in months}) == 1)
    annual, monthly = dc2["details"]["annual_fee"], dc2["details"]["monthly_revenue"]
    ok("DC2 the annual fee is 12 times the monthly revenue and under 12 times list", annual == 12 * monthly < 12 * dc2["details"]["list_price"])
    t = squash("\n".join(pdf_pages("04_customers/order_form_brookline_row.pdf")))
    ok("DC2 order form is an annual prepay at a discount to list", f"${annual:,}" in t and "once, in advance" in t and "off the" in t)
    ok("DC2 order form is a 12-month term with no pilot or exit-at-will clause", "12 months" in t and "pilot" not in t.lower() and "any reason" not in t)
    bill = {r["month"]: int(r["billings"]) for r in months}
    ok("DC2 billings show the whole annual fee in one month and nothing in the other months",
       bill[dc2["details"]["billings_month"]] == annual and sum(v for k, v in bill.items() if k != dc2["details"]["billings_month"]) == 0)
    ok("DC2 the billings month is a cash spike well above the clinic's revenue that month", bill[dc2["details"]["billings_month"]] > 10 * monthly)

    cap = workbook("02_financials/cap_table.xlsx")["Cap_Table"]
    rows = {cap.cell(row=r, column=1).value: r for r in range(1, 60) if isinstance(cap.cell(row=r, column=1).value, str)}
    r_tot = rows["Total"]
    holders = range(5, r_tot)           # the holders, then the Series A investors row with no shares yet
    total = cap.cell(row=r_tot, column=3).value
    ok("DC3 cap table holders add to the total", sum(cap.cell(row=r, column=3).value for r in holders) == total)
    ok("DC3 holder percentages add to 100.00 percent", near(sum(cap.cell(row=r, column=4).value for r in holders), 1.0, 1e-9))
    ok("DC3 pro forma percentages add to 100.00 percent", near(sum(cap.cell(row=r, column=6).value for r in holders), 1.0, 1e-9))
    safe_rows = range(rows["Investor"] + 1, rows["Total SAFEs"])
    amts = [cap.cell(row=r, column=2).value for r in safe_rows]
    caps = [cap.cell(row=r, column=3).value for r in safe_rows]
    shares = [cap.cell(row=r, column=6).value for r in safe_rows]
    ok("DC3 each SAFE's shares equal its amount over its cap price", all(near(sh, a * total / c, 0.5) for sh, a, c in zip(shares, amts, caps)))
    holder_safe = [cap.cell(row=r, column=3).value for r in holders if "SAFE converted" in str(cap.cell(row=r, column=1).value)]
    ok("DC3 the holder table carries the SAFE shares from the conversion table", holder_safe == shares)
    ok("DC3 the SAFEs own their amounts over their caps", near(sum(shares) / total, sum(a / c for a, c in zip(amts, caps)), 1e-6))
    ok("DC3 the SAFE amounts add to the money the key says was raised", sum(amts) == FACTS["raised_before_series_a"])
    pre, new = cap.cell(row=rows["Pre-money valuation (USD)"], column=2).value, cap.cell(row=rows["New money (USD)"], column=2).value
    ok("DC3 the Series A ownership is new money over post-money", near(cap.cell(row=rows["Series A ownership"], column=2).value, new / (pre + new), 1e-6))


# ---- the cash reconciles ------------------------------------------------------------------------
def check_cash():
    wb = workbook("02_financials/financial_model.xlsx")
    pl = wb["P&L"]
    op_row = next(r for r in range(1, 40) if str(pl.cell(row=r, column=1).value).startswith("Operating result"))
    loss = -sum(pl.cell(row=op_row, column=c).value for c in range(3, 24))
    raised = FACTS["raised_before_series_a"]
    pre = named(wb, "NET_LOSS_BEFORE_2025_USD")
    cash = named(wb, "CASH_START")
    ok("CASH raised less all losses equals cash at 2026-09-30, within a dollar", near(raised - pre - loss, cash, 1.0), f"{raised - pre - loss:,.2f} vs {cash:,}")
    ok("CASH the bridge sheet's difference is under a dollar", abs(named(wb, "CASH_BRIDGE_DIFF")) < 1.0)
    ok("CASH the loss before 2025 is a plausible 2024 spend", 100_000 <= pre <= 1_500_000, str(pre))
    revenue = sum(int(r["mrr"]) + int(r["one_time_fee"]) for r in cust_rows())
    spend = sum(float(r["amount_usd"]) for r in read_csv("02_financials/expense_ledger.csv"))
    ok("CASH the P&L losses equal the ledger spend less the customer file's revenue", near(loss, spend - revenue, 1.0), f"{loss:,.2f} vs {spend - revenue:,.2f}")


# ---- the room does not look built to be caught --------------------------------------------------
def check_realism():
    ids = [r["customer_id"] for r in cust_rows()]
    ok("every customer id follows the pattern C and four digits, the large customer and the pilot included", all(re.fullmatch(r"C\d{4}", i) for i in ids))
    ok("the ids of the large customer and the pilot are the ones the key names", FACTS["big_customer_id"] in ids and FACTS["pilot_customer_id"] in ids)
    deck13 = pdf_pages("01_overview/pitch_deck.pdf")[13]
    ok("deck slide 14 is titled Appendix: key metrics", "Appendix: key metrics" in deck13 and "Claims page" not in deck13)
    cover = squash("\n".join(pdf_pages("00_welcome/cover_note_and_index.pdf")))
    ok("cover note puts the second tier after a first meeting, with no term sheet", "on request after a first meeting" in cover and "term sheet" not in cover)
    ok("cover note does not promise every number ties", "ties to a file" not in cover and "every number" not in cover.lower())
    pl = workbook("02_financials/financial_model.xlsx")["P&L"]
    gm_row = next(r for r in range(1, 40) if pl.cell(row=r, column=1).value == "Gross margin")
    gms = [pl.cell(row=gm_row, column=c).value for c in range(3, 24)]
    ok("gross margin in the 21 actual months varies (at least 15 distinct values, spread over a point)",
       len({round(g, 4) for g in gms}) >= 15 and max(gms) - min(gms) > 0.01, f"{min(gms):.3f} to {max(gms):.3f}")
    pay_row = next(r for r in range(1, 40) if pl.cell(row=r, column=1).value == "Payroll")
    pay = [pl.cell(row=pay_row, column=c).value for c in range(3, 24)]
    steps = {round(b - a, 2) for a, b in zip(pay, pay[1:]) if b > a}
    ok("payroll does not rise by one fixed amount per hire", len(steps) >= 8, f"{len(steps)} distinct rises")
    ledger = read_csv("02_financials/expense_ledger.csv")
    ok("ledger has no vendor called Various and has more than 12 lines a month", all(r["vendor"] != "Various" for r in ledger) and len(ledger) > 21 * 12 + 60)
    rent = [float(r["amount_usd"]) for r in ledger if r["category"] == "rent"]
    ok("office rent moves in a few lease steps, not every month", len(set(rent)) <= 6, f"{len(set(rent))} values")
    prices = {int(r["mrr"]) for r in cust_rows() if r["month"] == "2026-09" and r["product_type"] == "Subscription"}
    ok("customer prices are not just the list prices (at least 15 distinct values)", len(prices) >= 15, str(len(prices)))
    locs = [int(r["locations"]) for r in cust_rows() if r["month"] == "2026-09" and r["product_type"] == "Subscription"]
    ok("some ordinary customers have more than one location", sum(1 for x in locs if x > 1) >= 10)
    big = next(r for r in cust_rows() if r["customer_id"] == FACTS["big_customer_id"] and r["month"] == "2026-09")
    pilot = next(r for r in cust_rows() if r["customer_id"] == FACTS["pilot_customer_id"] and r["month"] == "2026-09")
    ok("the large customer and the pilot are priced per location at or below Pro list ($1,099)",
       int(big["mrr"]) / int(big["locations"]) <= 1099 and int(pilot["mrr"]) / int(pilot["locations"]) < int(big["mrr"]) / int(big["locations"]))
    salaries = [int(r["base_salary_usd"]) for r in read_csv("05_team/hiring_plan.csv")]
    ok("hiring plan salaries are not all multiples of 5,000", sum(1 for x in salaries if x % 5000) >= 12)
    days = {r["start_date"][-2:] for r in read_csv("05_team/hiring_plan.csv")}
    ok("hire start dates are not all the first of the month", len(days) > 2)
    cap = workbook("02_financials/cap_table.xlsx")["Cap_Table"]
    total_row = next(r for r in range(1, 40) if cap.cell(row=r, column=1).value == "Total")
    total_shares = cap.cell(row=total_row, column=3).value
    ok("cap table share total is not a round number of millions", total_shares % 1_000_000 != 0, str(total_shares))
    bios = squash("\n".join(pdf_pages("05_team/founder_bios.pdf")))
    ok("the two founders do not share one GitHub handle", bios.count("github.com/") <= 1)


# ---- flags the room does not plant: the room must be honest on the point -----------------------
def deck_text():
    return pdf_pages("01_overview/pitch_deck.pdf")


def clean_pf1():
    wb = workbook("02_financials/financial_model.xlsx")
    rows = cust_rows()
    sub = sum(int(r["mrr"]) for r in rows if r["month"] == "2026-09" and r["product_type"] == "Subscription")
    ok("PF1 clean: model ARR equals 12 times subscription MRR, with the pilot left out", named(wb, "ARR_TOTAL") == 12 * sub)
    m = {r["month"]: r for r in read_csv("03_metrics/monthly_metrics.csv")}
    ok("PF1 clean: monthly_metrics arr (2026-09) leaves the pilot out", int(m["2026-09"]["arr"]) == 12 * sub)


def clean_pf2():
    wb = workbook("02_financials/financial_model.xlsx")
    ok("PF2 clean: deck ARR equals the model's ARR_TOTAL", num(claims()["ARR"][0]) == named(wb, "ARR_TOTAL"))


def clean_pf2b():
    text = (ROOM / "01_overview/pitch_video_script.md").read_text(encoding="utf-8")
    m = re.search(r"\*\*0:35[^\n]*\n\n(.*?)\n\n\*\*", text, re.S)
    para = m.group(1) if m else ""
    phrase = FACTS["arr_video_phrase"]
    deck = num(claims()["ARR"][0])
    ok("PF2b clean: the spoken ARR is a fair round-down of the deck's ARR", f"just over {phrase} dollars" in para and 0 <= deck - FACTS["arr_video"] < 500_000)


def clean_pf3():
    wb = workbook("02_financials/financial_model.xlsx")
    res = cohorts_from_csv()
    ts = sum(r["start"] for r in res)
    nrr = sum(r["end"] for r in res) / ts
    ok("PF3 clean: the Cohorts tab counts every customer", named(wb, "COHORT_CUSTOMERS") == sum(r["n"] for r in res))
    ok("PF3 clean: the Cohorts tab net retention equals the rebuilt figure", near(named(wb, "NRR_12M"), nrr, 0.001))
    ok("PF3 clean: the deck shows the same retention", f"{named(wb, 'NRR_12M') * 100:.1f}%" in deck_text()[5])


def clean_pf4():
    wb = workbook("02_financials/financial_model.xlsx")
    hires = hires_from_csv()
    runway, burn = rerun_cash(wb, hires)
    ok("PF4 clean: the model pays for every hire in the hiring plan", named(wb, "HIRES_COUNTED") == len(hires))
    ok("PF4 clean: model runway equals the rerun with all hires", runway == named(wb, "RUNWAY_MONTHS"), f"{runway} vs {named(wb, 'RUNWAY_MONTHS')}")
    ok("PF4 clean: the deck shows the model's runway and burn", claims()["Runway"][0] == f"{named(wb, 'RUNWAY_MONTHS')} months"
       and near(num(claims()["Average monthly net burn, next 12 months"][0]), burn, 1.0))


def clean_pf5():
    rows = [int(r["mrr"]) for r in cust_rows() if r["month"] == "2026-09" and r["product_type"] == "Subscription"]
    top1 = max(rows) / sum(rows)
    ok("PF5 clean: the deck states the largest customer's real share", claims()["Largest customer"][0] == f"{round(top1 * 100)}% of ARR")


def clean_pf6():
    wb = workbook("02_financials/financial_model.xlsx")
    led = defaultdict(float)
    for r in read_csv("02_financials/expense_ledger.csv"):
        if r["month"] == "2026-09":
            led[r["category"]] += float(r["amount_usd"])
    basis = int(next(r for r in read_csv("03_metrics/monthly_metrics.csv") if r["month"] == "2026-09")["mrr"])
    true = 1 - sum(led[k] for k in ("telephony", "third_party_licences", "cloud_hosting", "model_inference", "support_staff")) / basis
    ok("PF6 clean: model gross margin counts hosting, inference and support", near(named(wb, "GM_PCT"), true, 0.001), f"{named(wb, 'GM_PCT'):.4f} vs {true:.4f}")
    ws = wb["P&L"]
    labels = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 40)}
    ok("PF6 clean: hosting and support sit under cost of revenue", labels.get("Cloud hosting and model inference") == "COGS" and labels.get("Customer support specialists") == "COGS")
    ok("PF6 clean: the deck shows the true margin", claims()["Gross margin"][0] == f"{round(true * 100)}%")


def clean_pf7():
    paying = int(next(r for r in read_csv("03_metrics/monthly_metrics.csv") if r["month"] == "2026-09")["paying_accounts"])
    ok("PF7 clean: the deck clinic count is the paying count", num(claims()["Clinics on Hollis"][0]) == paying)


def check_key_shape():
    ok("the key's counts match its lists", KEY["counts"]["flags"] == len(KEY["flags"]) and KEY["counts"]["decoys"] == len(KEY["decoys"]) == 3)
    ok("the key's plants are the flags it lists", PLANTS == {f["id"] for f in KEY["flags"]} and set(KEY["plants"]) <= set(ALL_FLAGS))
    if KEY["mode"] == "main":
        ok("main key: 8 flags and 3 decoys", len(KEY["flags"]) == 8 and len(KEY["decoys"]) == 3)
        seed = {f["id"] for f in KEY["flags"] if f["catchable_from_seed_drop"]}
        ok("main key: five flags are catchable from the seed drop", seed == {"PF2", "PF2b", "PF4", "PF6", "PF7"})
    if KEY["mode"] == "clean":
        ok("clean key: no flags and it says so", not KEY["flags"] and KEY["clean_room"] is True)
    ok("every flag names a file location, a shown value, a true value and a check",
       all(f["locations"] and "value" in f["shown"] and "value" in f["true"] and f["check"] for f in KEY["flags"]))
    ok("every file a seed-drop flag needs is in the seed drop",
       all(x in KEY["seed_drop_files"] for f in KEY["flags"] if f["catchable_from_seed_drop"] for x in f["files_needed"]))
    ok("flags needing the full room name a file outside the seed drop",
       all(any(x not in KEY["seed_drop_files"] for x in f["files_needed"]) for f in KEY["flags"] if not f["catchable_from_seed_drop"]))


# ---- green flags: positives two different files agree on ---------------------------------------
# The pieces of the company as the readout groups them, and the planted flags each one holds
# (docs/ongoing/FRAMEWORKS.md, "The readout"). A piece with a planted flag reads weak, so a green flag
# can only come from a piece with none.
PIECE_FLAGS = {"traction and revenue": {"PF1", "PF2", "PF2b", "PF7"}, "retention": {"PF3"}, "burn and runway": {"PF4"},
               "unit economics": {"PF6"}, "customers": {"PF5"}, "market": set(), "team": set(), "the pitch itself": set()}
# The deck, its sidecar, the overview, the video script, the website and the one-pager and cover note repeat one another.
NARRATIVE = {"01_overview/pitch_deck.pdf", "01_overview/claims.json", "01_overview/company_overview.pdf",
             "01_overview/pitch_video_script.md", "06_product/website_copy.md", "06_product/product_one_pager.pdf",
             "00_welcome/cover_note_and_index.pdf"}


def read_location(loc):
    """The number a green-flag location holds, read from the room's file. Returns None when it cannot be found."""
    rel = loc["file"]
    if rel.endswith(".xlsx"):
        wb = workbook(rel)
        val = wb[loc["sheet"]][loc["cell"]].value
        if "named_cell" in loc and named(wb, loc["named_cell"]) != val:
            return None
        return val if isinstance(val, (int, float)) else None
    if rel.endswith(".pdf"):
        if squash(loc["text"]) not in squash(pdf_pages(rel)[loc["page"] - 1]):
            return None
        nums = re.findall(r"\d[\d,]*(?:\.\d+)?", loc["text"])      # the last number in the quoted text
        return float(nums[-1].replace(",", "")) if nums else None
    if rel.endswith(".csv"):
        rows = [r for r in read_csv(rel) if all(r.get(k) == v for k, v in loc["row_key"].items())]
        return num(rows[0][loc["column"]]) if len(rows) == 1 else None
    return None


def independent_count(locations):
    files = {l["file"] for l in locations}
    return len(files - NARRATIVE) + (1 if files & NARRATIVE else 0)


def check_one_green(gf, entry, parent):
    tag = f"{gf['id']} ({gf['piece']}, {entry['label']} {entry['number']:,})"
    locs = entry["locations"]
    ok(f"{tag}: two or more locations in the room", len(locs) >= 2)
    vals = [read_location(l) for l in locs]
    ok(f"{tag}: every location is found in its file", all(v is not None for v in vals), str(vals))
    ok(f"{tag}: every location agrees with the number within {parent['tolerance'] * 100:.0f} percent",
       all(v is not None and near(v, entry["number"], 0, parent["tolerance"]) for v in vals), str(vals))
    ok(f"{tag}: the locations are in at least two independent sources (the deck, overview, video and website count as one)",
       independent_count(locs) >= 2, str(independent_count(locs)))
    if "independent_sources" in entry:
        ok(f"{tag}: the key's independent source count is right", independent_count(locs) == entry["independent_sources"])


def check_green_flags():
    gfs = KEY.get("green_flags")
    ok("the key has a green_flags list", bool(gfs))
    if not gfs:
        return
    ok("the key's green flags are one per strong piece, three in all", len(gfs) == KEY["counts"]["green_flags"] == 3
       and {g["piece"] for g in gfs} == {"the pitch itself", "market", "team"})
    for gf in gfs:
        flags_in_piece = PIECE_FLAGS.get(gf["piece"])
        ok(f"{gf['id']}: {gf['piece']} is a piece with no planted flag", flags_in_piece is not None and not flags_in_piece)
        ok(f"{gf['id']}: no planted flag in this room belongs to {gf['piece']}", not (PLANTS & (flags_in_piece or set())))
        check_one_green(gf, gf, gf)
        for extra in gf.get("also_agree", []):
            check_one_green(gf, extra, gf)


# What the main room's planted flags must read. The values are the ones the answer key held before the green
# flags were added. Any change to a shown or true number, or to a flag's supporting figures, fails here.
RED_BASELINE = {
    "PF1": (2240868, 2060868), "PF2": (2500000, 2240868), "PF2b": (3000000, 2500000), "PF3": (1.026, 0.873),
    "PF4": (37, 19), "PF5": (0.09, 0.21), "PF6": (0.8, 0.581), "PF7": (357, 212),
}
RED_DETAILS = {
    "PF1": ({"pilot_arr_counted": 180000, "metrics_arr_2026_09": 2240868}, {"pilot_fees_contracted_total": 90000, "pilot_notice_days": 30}),
    "PF2": ({"gap_to_model": 259132}, {"recurring_arr": 2060868}),
    "PF3": ({"customers": 57}, {"gross_revenue_retention": 0.734, "customers": 68, "customers_dropped": 11}),
    "PF4": ({"hires_counted": 8, "avg_net_burn_first_12_months": 374462}, {"hires_in_plan": 24, "avg_net_burn_first_12_months": 480062}),
    "PF5": (None, {"top3_share": 0.222, "largest_customer_id": "C0087"}),
}


def check_red_flags_unchanged():
    if KEY["mode"] != "main":
        ok("red flags: this is not the main room, so the pinned numbers do not apply; each flag is still checked against its own key", True)
        return
    ok("red flags: the main room plants all eight flags", PLANTS == set(ALL_FLAGS))
    for fid, (shown, true) in RED_BASELINE.items():
        f = flag(fid)
        ok(f"red flag {fid}: shown {shown} and true {true} are unchanged", f["shown"]["value"] == shown and f["true"]["value"] == true,
           f"{f['shown']['value']} / {f['true']['value']}")
    for fid, (sd, td) in RED_DETAILS.items():
        f = flag(fid)
        ok(f"red flag {fid}: the supporting figures are unchanged",
           all(f["shown"].get("details", {}).get(k) == v for k, v in (sd or {}).items())
           and all(f["true"].get("details", {}).get(k) == v for k, v in td.items()))


def check_team_piece():
    team = FACTS["team"]
    deck10 = squash(pdf_pages("01_overview/pitch_deck.pdf")[9])
    bios = squash("\n".join(pdf_pages("05_team/founder_bios.pdf")))
    overview = squash("\n".join(pdf_pages("01_overview/company_overview.pdf")))
    cap = workbook("02_financials/cap_table.xlsx")["Cap_Table"]
    cap_names = " ".join(str(cap.cell(row=r, column=1).value) for r in range(1, 20))
    senior = team[2:]
    ok("team: every senior name on deck slide 10 is in the bios under the same name", all(t["name"] in deck10 and t["name"] in bios for t in senior))
    for t in team:
        m = re.search(re.escape(t["name"]) + r"\s+" + re.escape(t["title"]) + r"\s+(?:GitHub: \S+\s+)?(\w+)", bios)
        ok(f"team: the bio of {t['name']} uses the first name {t['name'].split()[0]}", bool(m) and m.group(1) == t["name"].split()[0],
           m.group(1) if m else "no bio found")
    ok("team: the founders' names agree across the deck, bios, overview, video script and cap table",
       all(t["name"] in deck10 and t["name"] in bios and t["name"] in overview and t["name"] in cap_names for t in team[:2])
       and all(t["name"] in (ROOM / "01_overview/pitch_video_script.md").read_text(encoding="utf-8") for t in team[:1]))
    plan = read_csv("05_team/hiring_plan.csv")
    ok("team: no planned hire has the title of a current senior person", not ({r["role"] for r in plan} & {t["title"] for t in senior}))
    after = [int(r["headcount_after_hire"]) for r in plan]
    ok("team: the hiring plan's staff count rises by one a hire, from today's staff to the deck's staff after hiring",
       after == list(range(FACTS["staff_now"] + 1, FACTS["staff_after_hiring"] + 1)) and len(plan) == 24)
    wb = workbook("02_financials/financial_model.xlsx")
    stray = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.lower().startswith("staff after")]
    ok("team: the model states no staff total after the hires", not stray, str(stray))
    pl = wb["P&L"]
    heads = next(r for r in range(1, 40) if pl.cell(row=r, column=1).value == "Staff at month end")
    ok("team: the P&L shows staff for the 21 actual months only, ending at today's staff",
       all(pl.cell(row=heads, column=c).value is not None for c in range(3, 24)) and pl.cell(row=heads, column=23).value == FACTS["staff_now"]
       and all(pl.cell(row=heads, column=c).value is None for c in range(24, 36)))
    ok("team: deck slide 10, the overview and the model all say staff today is 17",
       f"{FACTS['staff_now']} people today" in deck10 and f"{FACTS['staff_now']} people work there today" in overview
       and named(wb, "STAFF_NOW") == FACTS["staff_now"])


def check_market_piece():
    wb = workbook("02_financials/financial_model.xlsx")
    ok("market: the model's Market tab multiplies clinics by annual value to get the market size",
       named(wb, "MARKET_TAM_USD") == named(wb, "MARKET_CLINICS") * named(wb, "MARKET_ACV_USD"))
    got = claims()
    ok("market: the deck's clinic count and market size match the Market tab",
       num(got["US independent clinics (company estimate)"][0]) == named(wb, "MARKET_CLINICS")
       and num(got["Market size (company estimate)"][0]) == named(wb, "MARKET_TAM_USD"))
    note = squash(str(wb["Market"]["A2"].value))
    ok("market: the Market tab says the figures are the company's own estimates", "company's own estimates" in note)


FLAG_CHECKS = {"PF1": check_pf1, "PF2": check_pf2, "PF2b": check_pf2b, "PF3": check_pf3, "PF4": check_pf4,
               "PF5": check_pf5, "PF6": check_pf6, "PF7": check_pf7}
CLEAN_CHECKS = {"PF1": clean_pf1, "PF2": clean_pf2, "PF2b": clean_pf2b, "PF3": clean_pf3, "PF4": clean_pf4,
                "PF5": clean_pf5, "PF6": clean_pf6, "PF7": clean_pf7}


def main():
    print(f"Checking {ROOM} against the key with plants: {', '.join(sorted(PLANTS)) or 'none'}")
    check_files()
    check_formulas_have_values()
    check_locations()
    check_claims_page()
    for fid in ALL_FLAGS:
        (FLAG_CHECKS if fid in PLANTS else CLEAN_CHECKS)[fid]()
    check_decoys()
    check_cash()
    check_realism()
    check_key_shape()
    check_green_flags()
    check_red_flags_unchanged()
    check_team_piece()
    check_market_piece()
    print(f"\n{len(PASSES)} passed, {len(FAILS)} failed")
    if FAILS:
        print("FAILED: " + "; ".join(FAILS))
        sys.exit(1)
    if PLANTS:
        print("All planted flags and decoys are in the room exactly as the answer key says.")
    else:
        print("The room is clean: no flag is planted, and every number a scan could test is honest.")


if __name__ == "__main__":
    main()
