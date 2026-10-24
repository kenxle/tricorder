"""Chart data for the expert-read report sections, computed by script from the room's files.

    traction_chart(readout, room_dir) -> dict | None
    burn_chart(readout, room_dir) -> dict | None

Each chart is plain JSON a page can draw: a title, a unit, series (each with x,
y, the file it came from and the locator inside that file) and markers (dates
or figures to mark). Nothing here comes from a model. A missing file means no
chart, never an error.

Traction and revenue: the metrics file's monthly figure (ARR when the deck
states ARR, else MRR), the same figure rebuilt from the customer file's
subscription rows, and the pitch's stated figure marked against both.

Burn and runway: cash by month under the company's own plan (the model's
Cash_Runway tab), and the same cash roll with the burn the checks found that
the plan leaves out: the hires in hiring_plan.csv the model does not count
(FRAMEWORKS 1.4) and the contractor spend the plan drops (1.12 kind e). When
the checks found nothing, the second line holds the last actual month's net
burn flat. Each line's runway end is marked, with Kimmy's baseline of runway
after the round closes, read from her knowledge file.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path

from schema.claim import Claim, Readout

from scans.knowledge import ROOT as KNOWLEDGE_ROOT

_MONTH = re.compile(r"^\d{4}-\d{2}$")


# Small helpers ---------------------------------------------------------------


def add_months(month: str, n: int) -> str:
    y, m = int(month[:4]), int(month[5:7])
    k = y * 12 + (m - 1) + n
    return f"{k // 12:04d}-{k % 12 + 1:02d}"


def months_between(a: str, b: str) -> int:
    return (int(b[:4]) * 12 + int(b[5:7])) - (int(a[:4]) * 12 + int(a[5:7]))


def _rel(path: Path, room: Path) -> str:
    try:
        return str(path.relative_to(room))
    except ValueError:
        return path.name


def _num(v) -> float | None:
    try:
        return float(str(v).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def _read_csv(path: Path) -> list[dict[str, str]]:
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _metrics_file(room: Path) -> Path | None:
    hits = sorted(room.glob("*/monthly_metrics.csv"))
    return hits[0] if hits else None


def _customer_file(room: Path) -> Path | None:
    for p in sorted(room.glob("03_metrics/*.csv")):
        if p.name == "monthly_metrics.csv":
            continue
        with open(p, newline="", encoding="utf-8") as fh:
            head = next(csv.reader(fh), [])
        if "customer_id" in head and "month" in head:
            return p
    return None


def _is_pitch(claim: Claim) -> bool:
    return Path(claim.sources[0].uri).name.startswith("pitch_")


def _pitch_words(uri: str) -> str:
    name = Path(uri).name
    if name.startswith("pitch_deck"):
        return "Deck"
    if name.startswith("pitch_video"):
        return "Pitch video"
    return name


# Traction and revenue --------------------------------------------------------

# The metrics column for each figure the pitch states, in the order tried.
_METRIC_COLUMNS = {"arr": ("arr",), "mrr": ("mrr", "subscription_revenue")}
_UNITS = {"arr": "USD_per_year", "mrr": "USD_per_month"}


def stated_figures(readout: Readout) -> list[Claim]:
    """The pitch's ARR or MRR claims, one per place it says it, deck first."""
    seen, out = set(), []
    for c in readout.claims:
        if c.value and c.value.label in _METRIC_COLUMNS and _is_pitch(c):
            key = (c.sources[0].uri, c.sources[0].locator, round(c.value.number, 2))
            if key in seen:
                continue
            seen.add(key)
            out.append(c)
    return sorted(out, key=lambda c: (not Path(c.sources[0].uri).name.startswith("pitch_deck"), c.sources[0].uri))


def traction_chart(readout: Readout, room_dir: str | Path) -> dict | None:
    room = Path(room_dir)
    metrics = _metrics_file(room)
    if metrics is None:
        return None
    rows = _read_csv(metrics)
    if not rows or "month" not in rows[0]:
        return None
    stated = stated_figures(readout)
    label = next((c.value.label for c in stated), "mrr")
    stated = [c for c in stated if c.value.label == label]
    column = next((col for col in _METRIC_COLUMNS[label] if col in rows[0]), None)
    if column is None:
        label, stated = "mrr", []
        column = next((col for col in ("mrr", "subscription_revenue", "total_revenue") if col in rows[0]), None)
        if column is None:
            return None
    word = label.upper()
    pts = [(r["month"], _num(r[column])) for r in rows if _MONTH.match(r.get("month", "")) and _num(r[column]) is not None]
    if not pts:
        return None
    series = [{
        "name": f"{word} in the metrics file",
        "role": "metric",
        "x": [m for m, _ in pts], "y": [v for _, v in pts],
        "source": {"file": _rel(metrics, room), "locator": f"column {column}, rows 2 to {len(rows) + 1}"},
    }]
    customers = _customer_file(room)
    if customers is not None:
        crow = _read_csv(customers)
        head = crow[0].keys() if crow else []
        col = "mrr" if "mrr" in head else ("subscription_revenue" if "subscription_revenue" in head else None)
        if col:
            filt = "product_type" in head
            sums: dict[str, float] = {}
            for r in crow:
                if filt and r.get("product_type") != "Subscription":
                    continue
                v = _num(r.get(col))
                if v is not None and _MONTH.match(r.get("month", "")):
                    sums[r["month"]] = sums.get(r["month"], 0.0) + v
            if sums:
                times = 12 if label == "arr" else 1
                xs = sorted(sums)
                where = f"column {col}" + (", rows with product_type Subscription" if filt else "") + ", summed by month"
                series.append({
                    "name": f"{word} rebuilt from the customer file's subscription rows" + (" (12 x the month)" if times == 12 else ""),
                    "role": "rebuilt",
                    "x": xs, "y": [round(sums[m] * times, 2) for m in xs],
                    "source": {"file": _rel(customers, room), "locator": where},
                })
    markers = []
    for c in stated:
        s = c.sources[0]
        markers.append({
            "name": f"{_pitch_words(s.uri)}, {s.locator}",
            "role": "stated",
            "x": c.value.period or series[0]["x"][-1], "y": c.value.number,
            "source": {"file": s.uri, "locator": s.locator},
            "claim_ids": [c.id],
        })
    return {
        "piece": "traction_and_revenue",
        "kind": "revenue_over_time",
        "title": f"{word} by month, against what the pitch states",
        "unit": _UNITS[label],
        "series": series,
        "markers": markers,
        "lines": [],
    }


# Burn and runway -------------------------------------------------------------


def _workbook(room: Path):
    hits = sorted(room.glob("*/financial_model.xlsx"))
    if not hits:
        return None, None
    try:
        import openpyxl
        return openpyxl.load_workbook(hits[0], data_only=True), hits[0]
    except Exception:  # an unreadable workbook is no chart, not an error
        return None, None


def _named(wb, name: str):
    dn = wb.defined_names.get(name)
    if dn is None:
        return None
    for sheet, ref in dn.destinations:
        try:
            return wb[sheet][ref.replace("$", "")].value
        except KeyError:
            return None
    return None


def _cash_plan(wb) -> dict | None:
    """Cash_Runway: month, equity raised, net burn and ending cash, with the cells they come from."""
    if "Cash_Runway" not in wb.sheetnames:
        return None
    ws = wb["Cash_Runway"]
    head_row, cols = None, {}
    for row in ws.iter_rows():
        names = {str(c.value).strip().lower(): c.column_letter for c in row if c.value is not None}
        if "month" in names and "ending cash" in names:
            head_row, cols = row[0].row, names
            break
    if head_row is None:
        return None
    out = {"months": [], "cash": [], "equity": [], "burn": [], "first": None, "last": None, "cols": cols}
    for r in range(head_row + 1, ws.max_row + 1):
        m = ws[f"{cols['month']}{r}"].value
        if not (isinstance(m, str) and _MONTH.match(m)):
            if out["months"]:
                break
            continue
        cash = _num(ws[f"{cols['ending cash']}{r}"].value)
        if cash is None:
            continue
        out["months"].append(m)
        out["cash"].append(cash)
        out["equity"].append(_num(ws[f"{cols['equity raised']}{r}"].value) or 0.0 if "equity raised" in cols else 0.0)
        out["burn"].append(_num(ws[f"{cols['net burn']}{r}"].value) or 0.0 if "net burn" in cols else 0.0)
        out["first"] = out["first"] or r
        out["last"] = r
    return out if out["months"] else None


def runway_end(months: list[str], cash: list[float]) -> str | None:
    """The last month with cash above zero before cash first reaches zero; None when it never does."""
    for i, c in enumerate(cash):
        if c <= 0:
            return months[i - 1] if i > 0 else None
    return None


def _hiring_rows(wb) -> list[tuple[str, str]]:
    """(role, start month) for each hire the model's hiring tab pays for."""
    out = []
    for name in ("Hiring_Plan", "Team"):
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        on = False
        for row in ws.iter_rows(values_only=True):
            a = str(row[0]).strip() if row and row[0] is not None else ""
            if a.lower() in ("role", "planned hire"):
                on = True
                continue
            if on:
                if not a or a.lower().startswith(("hires", "employees", "synthetic")):
                    on = False
                    continue
                start = str(row[2])[:7] if len(row) > 2 and row[2] is not None else ""
                out.append((a.lower(), start))
    return out


def missing_hires(wb, room: Path) -> tuple[list[dict], Path | None]:
    """Hires in hiring_plan.csv that the model's hiring tab does not count, matched by role and start month."""
    plans = sorted(room.glob("*/hiring_plan.csv"))
    if not plans:
        return [], None
    counted = Counter(_hiring_rows(wb))
    out = []
    for i, r in enumerate(_read_csv(plans[0]), start=2):
        key = ((r.get("role") or "").strip().lower(), (r.get("start_date") or "")[:7])
        if counted[key] > 0:
            counted[key] -= 1
            continue
        base = _num(r.get("base_salary_usd"))
        if base is None or not _MONTH.match(key[1]):
            continue
        out.append({"row": i, "id": r.get("hire_id", ""), "role": r.get("role", ""), "start": key[1], "base": base})
    return out, plans[0]


def _pnl_row_gap(wb, locator: str) -> tuple[str, int, dict[str, float]] | None:
    """For a P&L cell such as 'P&L!S23': the row's line name, and per plan month how far the plan
    sits below the last actual month."""
    m = re.match(r"^'?([^'!]+)'?!\$?[A-Z]+\$?(\d+)$", locator.strip())
    if not m or m.group(1) not in wb.sheetnames:
        return None
    ws, row = wb[m.group(1)], int(m.group(2))
    head = kind = None
    for r in range(1, 12):
        if ws.cell(r, 1).value == "Line":
            head, kind = r, r + 1
            break
    if head is None:
        return None
    last_actual, plan = None, {}
    for c in range(3, ws.max_column + 1):
        month, k = ws.cell(head, c).value, str(ws.cell(kind, c).value or "")
        v = _num(ws.cell(row, c).value)
        if not (isinstance(month, str) and _MONTH.match(month)) or v is None:
            continue
        if k.lower().startswith("actual"):
            last_actual = v
        elif k.lower().startswith("plan"):
            plan[month] = v
    if last_actual is None or not plan:
        return None
    return str(ws.cell(row, 1).value), row, {mo: round(last_actual - v, 2) for mo, v in plan.items() if last_actual - v > 0}


def kimmy_runway_baseline(root: Path = KNOWLEDGE_ROOT) -> int | None:
    """The months of runway after the round closes that Kimmy's burn and runway file gives as a baseline."""
    try:
        text = (root / "burn_and_runway.md").read_text(encoding="utf-8")
    except OSError:
        return None
    m = re.search(r"(\d+) months of runway after the round closes", text)
    return int(m.group(1)) if m else None


def burn_chart(readout: Readout, room_dir: str | Path) -> dict | None:
    room = Path(room_dir)
    wb, path = _workbook(room)
    if wb is None:
        return None
    plan = _cash_plan(wb)
    if plan is None:
        return None
    months, cash = plan["months"], plan["cash"]
    rel = _rel(path, room)
    cols = plan["cols"]
    cash_loc = f"Cash_Runway!{cols['ending cash']}{plan['first']}:{cols['ending cash']}{plan['last']}"
    series = [{"name": "Cash under the company's plan", "role": "plan", "x": months, "y": [round(c, 2) for c in cash],
               "source": {"file": rel, "locator": cash_loc}}]

    # What the checks found the plan leaves out, as extra spend per month.
    extra = [0.0] * len(months)
    found, sources, claim_ids = [], [{"file": rel, "locator": cash_loc}], []
    hires, plan_file = missing_hires(wb, room)
    load = _num(_named(wb, "LOAD_FACTOR"))
    if hires and load:
        for h in hires:
            cost = h["base"] * load / 12
            for i, mo in enumerate(months):
                if mo >= h["start"] and i > 0:
                    extra[i] += cost
        rows = f"rows {hires[0]['row']} to {hires[-1]['row']}" if len(hires) > 1 else f"row {hires[0]['row']}"
        found.append(f"{len(hires)} hires in the hiring plan that the model does not pay for, at the model's load factor")
        sources.append({"file": _rel(plan_file, room), "locator": f"{rows}, base_salary_usd times LOAD_FACTOR, divided by 12"})
        claim_ids += [c.id for c in readout.claims if c.value and c.value.label == "hires" and c.severity == "flag"]
    for c in readout.claims:
        if c.check.name != "contractor_spend_vs_headcount_and_burn" or c.severity != "flag":
            continue
        gap = _pnl_row_gap(wb, c.sources[0].locator)
        if not gap:
            continue
        line, row, by_month = gap
        last_gap = list(by_month.values())[-1] if by_month else 0.0
        for i, mo in enumerate(months):
            if i > 0:
                extra[i] += by_month.get(mo, last_gap if mo > max(by_month) else 0.0)
        found.append(f'the "{line}" spend the plan drops after the last actual month')
        sources.append({"file": rel, "locator": f"P&L row {row}, last actual month less each plan month"})
        claim_ids.append(c.id)

    if found:
        name = "Cash with the burn the checks found"
        roll, spent = [], 0.0
        for c, e in zip(cash, extra):
            spent += e
            roll.append(round(c - spent, 2))
        what = "The company's plan, plus " + " and ".join(found) + "."
    else:
        # Nothing found: hold the last actual month's net burn flat, with the plan's own raise.
        burn = _last_actual_burn(wb)
        if burn is None:
            return _burn_result(series, months, cash, None, plan, wb, rel)
        name = "Cash at the last actual month's net burn, held flat"
        roll = [round(cash[0], 2)]
        for i in range(1, len(months)):
            roll.append(round(roll[-1] + plan["equity"][i] - burn[0], 2))
        what = f"The plan's own raise, with net burn held at {burn[1]}'s level."
        sources.append({"file": rel, "locator": f"P&L row {burn[2]}, {burn[1]}"})
    series.append({"name": name, "role": "found", "x": months, "y": roll, "source": sources[1] if len(sources) > 1 else sources[0],
                   "sources": sources, "claim_ids": sorted(set(claim_ids)), "what": what})
    return _burn_result(series, months, cash, roll, plan, wb, rel)


def _last_actual_burn(wb) -> tuple[float, str, int] | None:
    if "P&L" not in wb.sheetnames:
        return None
    ws = wb["P&L"]
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").lower().startswith("net burn"):
            last = None
            for c in range(3, ws.max_column + 1):
                if str(ws.cell(5, c).value or "").lower().startswith("actual") and _num(ws.cell(r, c).value) is not None:
                    last = (_num(ws.cell(r, c).value), str(ws.cell(4, c).value), r)
            return last
    return None


def _burn_result(series, months, cash, roll, plan, wb, rel) -> dict:
    lines = []
    plan_end = runway_end(months, cash)
    lines.append({"name": "Runway ends, company's plan", "role": "plan", "x": plan_end,
                  "months": months_between(months[0], plan_end) if plan_end else None,
                  "beyond": plan_end is None, "source": series[0]["source"]})
    if roll is not None:
        found_end = runway_end(months, roll)
        lines.append({"name": "Runway ends, with the burn found", "role": "found", "x": found_end,
                      "months": months_between(months[0], found_end) if found_end else None,
                      "beyond": found_end is None, "source": series[1]["source"]})
    base = kimmy_runway_baseline()
    close = next((months[i] for i, e in enumerate(plan["equity"]) if i > 0 and e > 0), None)
    if base and close:
        cols = plan["cols"]
        lines.append({"name": f"Kimmy's baseline: {base} months after the round closes", "role": "baseline",
                      "x": add_months(close, base), "months": base, "close": close,
                      "source": {"file": "knowledge/burn_and_runway.md", "locator": "Stage lens, Runway; the close month from "
                                 f"Cash_Runway column {cols.get('equity raised', '?')} in {rel}"}})
    stated = _named(wb, "RUNWAY_MONTHS")
    return {
        "piece": "burn_and_runway",
        "kind": "cash_and_runway",
        "title": "Cash by month, and when it runs out",
        "unit": "USD",
        "data_month": months[0],
        "model_runway_months": stated,
        "series": [s for s in series if s],
        "markers": [],
        "lines": [ln for ln in lines if ln["x"]],
        "beyond": [ln["name"] for ln in lines if not ln["x"]],
    }
