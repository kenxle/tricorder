"""Accounting mislabels, FRAMEWORKS.md entry 1.12.

Each test names the requirement it proves (docs/features/20261022.12_accounting_mislabels/03_prd.md).
Figures in the asserts come from the spike run in that folder (spike/spike_run.txt).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import socket
from pathlib import Path

import pytest

from eval.score import score
from schema.check_quotes import check_claim
from schema.claim import Claim, Gap, Readout
from scans.base import Ids
from scans.redflags import mislabels

REPO = Path(__file__).resolve().parents[3]
ROOMS = REPO / "data_room"
FULL = ROOMS / "hollis_voice"
SEED = ROOMS / "hollis_voice_seed_drop"
CLEAN = ROOMS / "hollis_voice_clean"
BRASS = ROOMS / "brasswick"
ALL_ROOMS = [FULL, SEED, CLEAN, BRASS]

A, B, C = "mrr_column_vs_subscription_rows", "billings_vs_recognised_revenue", "gross_vs_net_revenue"
E, F, G, H = ("contractor_spend_vs_headcount_and_burn", "capitalised_spend_vs_pnl_and_cash",
              "deferred_revenue_vs_earned", "one_time_fees_in_recurring_metrics")
ALL_CHECKS = [A, B, C, E, F, G, H]
CUST = "03_metrics/customer_mrr_by_month.csv"
METRICS = "03_metrics/monthly_metrics.csv"


_CACHE: dict[str, tuple[list[Claim], list[Gap]]] = {}


def chk(room: Path) -> tuple[list[Claim], list[Gap]]:
    claims, gaps, _panels = mislabels.check(room, Ids("redflags"))
    return claims, gaps


def run(room: Path) -> tuple[list[Claim], list[Gap]]:
    key = str(room)
    if key not in _CACHE:
        _CACHE[key] = chk(room)
    return _CACHE[key]


def by_check(claims, name, severity=None):
    return [c for c in claims if c.check.name == name and (severity is None or c.severity == severity)]


def cites(claim, uri, locator=None):
    return any(s.uri == uri and (locator is None or s.locator == locator) for s in claim.sources)


def outcomes(room):
    return mislabels.outcomes(room)


# Test rooms with a resale line (AM-R30). Built in tmp_path; nothing under data_room/ changes.
def _copy(dst: Path) -> Path:
    shutil.copytree(FULL, dst)
    return dst


@pytest.fixture(scope="module")
def resale_rooms(tmp_path_factory):
    base = tmp_path_factory.mktemp("resale")
    off = _copy(base / "hollis_voice_resale_off")
    on = _copy(base / "hollis_voice_resale_on")
    lines = (FULL / CUST).read_text(encoding="utf-8").splitlines()
    rows = [next(csv.reader([ln])) for ln in lines[1:]]
    out, extra = [], {}
    for r in rows:
        out.append(r)
        if r[5] == "Subscription" and r[0] >= "2026-04":
            amt = 25 * int(r[3])
            out.append([*r[:5], "Resale", str(amt), "0", str(amt)])
            extra[r[0]] = extra.get(r[0], 0) + amt
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(lines[0].split(","))
    w.writerows(out)
    (on / CUST).write_text(buf.getvalue(), encoding="utf-8")
    met = (FULL / METRICS).read_text(encoding="utf-8").splitlines()
    new = [met[0]]
    for ln in met[1:]:
        f = ln.split(",")
        if f[0] in extra:
            mrr = int(f[1]) + extra[f[0]]
            f[1], f[2] = str(mrr), str(12 * mrr)
        new.append(",".join(f))
    (on / METRICS).write_text("\n".join(new) + "\n", encoding="utf-8")
    return off, on


def _tree_hash(root: Path) -> str:
    h = hashlib.sha256()
    for p in sorted(root.rglob("*")):
        if p.is_file() and ".venv" not in p.parts and "__pycache__" not in p.parts:
            h.update(str(p.relative_to(root)).encode())
            h.update(p.read_bytes())
    return h.hexdigest()


# ---------------------------------------------------------------- shared


def test_check_returns_claims_and_gaps():  # AM-R1
    claims, gaps = run(FULL)
    assert claims and all(isinstance(c, Claim) for c in claims)
    assert gaps and all(isinstance(g, Gap) for g in gaps)


def test_runs_with_network_blocked(monkeypatch):  # AM-R1
    def refuse(*a, **k):
        raise AssertionError("mislabels made a network call")
    monkeypatch.setattr(socket.socket, "connect", refuse)
    chk(BRASS)


def test_makes_no_model_call(monkeypatch):  # AM-R1, AM-R10: the fallback reads the model file, never an LLM
    import clients.llm as llm
    for name in dir(llm):
        obj = getattr(llm, name)
        if callable(obj) and not isinstance(obj, type) and getattr(obj, "__module__", "") == llm.__name__:
            monkeypatch.setattr(llm, name, lambda *a, **k: (_ for _ in ()).throw(AssertionError("LLM called")))
    chk(SEED)


def test_tolerance_override_from_settings(tmp_path, monkeypatch):  # AM-R1
    p = tmp_path / "settings.json"
    p.write_text(json.dumps({"match_tolerance": 0.10}), encoding="utf-8")
    monkeypatch.setenv("TRICORDER_SETTINGS", str(p))
    claims, _ = chk(FULL)
    assert not by_check(claims, A, "flag")  # 8.7% is under 10%


@pytest.mark.parametrize("room", ALL_ROOMS, ids=lambda p: p.name)
def test_every_claim_is_script_computed_with_its_check_name(room):  # AM-R2
    for c in run(room)[0]:
        assert c.check.kind == "script" and c.check.model is None
        assert c.check.name in ALL_CHECKS
        assert c.confidence == "computed"
        assert c.sources and c.conflicts_with == [] and c.corroborated_by == []


@pytest.mark.parametrize("room", ALL_ROOMS, ids=lambda p: p.name)
def test_flag_text_names_the_mislabel_kind(room):  # AM-R3
    for c in run(room)[0]:
        if c.severity in ("flag", "watch"):
            assert "Mislabel kind " in c.text, c.text


@pytest.mark.parametrize("room", ALL_ROOMS, ids=lambda p: p.name)
def test_gaps_use_entry_1_12_wording_and_piece(room):  # AM-R4
    for g in run(room)[1]:
        assert g.what.startswith("Not in the room:")
        assert g.piece in ("traction_and_revenue", "burn_and_runway")
        assert g.would_need


@pytest.mark.parametrize("room", ALL_ROOMS, ids=lambda p: p.name)
def test_every_claim_passes_check_quotes(room):  # AM-R5
    for c in run(room)[0]:
        _, problems = check_claim(c, room)
        assert not problems, (c.id, problems)


def test_every_claim_passes_check_quotes_in_resale_rooms(resale_rooms):  # AM-R5
    for room in resale_rooms:
        for c in chk(room)[0]:
            assert not check_claim(c, room)[1]


def test_missing_column_is_a_gap_not_a_crash(tmp_path):  # AM-R6
    room = _copy(tmp_path / "room")
    lines = (room / CUST).read_text(encoding="utf-8").splitlines()
    (room / CUST).write_text("\n".join(",".join(ln.split(",")[:8]) for ln in lines) + "\n", encoding="utf-8")
    claims, gaps = chk(room)
    assert not by_check(claims, B)
    assert any("billings" in g.what.lower() for g in gaps)
    assert by_check(claims, A, "flag")  # the others still run


def test_missing_pnl_row_is_a_gap(tmp_path):  # AM-R6
    import openpyxl
    room = _copy(tmp_path / "room")
    path = room / "02_financials/financial_model.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["P&L"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "Subscription revenue":
            ws.cell(r, 1).value = "Renamed row"
    wb.save(path)
    claims, gaps = chk(room)
    assert not by_check(claims, B, "flag")
    assert any("Subscription revenue" in g.what for g in gaps)


def test_locator_error_is_caught_per_subcheck(monkeypatch):  # AM-R6
    from schema.readers import LocatorError

    def boom(*a, **k):
        raise LocatorError("broken")
    monkeypatch.setattr(mislabels, "_check_e", boom)
    claims, gaps = chk(FULL)
    assert any("broken" in g.why for g in gaps)
    assert by_check(claims, A, "flag")


# ---------------------------------------------------------------- kind a


def test_full_room_mrr_column_holds_the_pilot():  # AM-R7, AM-R8
    flags = by_check(run(FULL)[0], A, "flag")
    assert len(flags) == 1
    c = flags[0]
    assert c.value.number == 171739 and c.value.label == "mrr" and c.value.period == "2026-09"
    assert "$186,739" in c.text and "$15,000" in c.text and "8.7%" in c.text
    assert "Marlow Veterinary Group" in c.text and "2026-08" in c.text
    assert cites(c, CUST, "row 2101") and cites(c, METRICS)


def test_mrr_claim_cites_the_pilot_arr_build_row():  # AM-R9
    c = by_check(run(FULL)[0], A, "flag")[0]
    assert cites(c, "02_financials/financial_model.xlsx", "ARR_Build!E13")
    assert cites(c, "02_financials/financial_model.xlsx", "ARR_Build!A13")


def test_seed_drop_mrr_uses_arr_build_fallback():  # AM-R10
    claims, gaps = run(SEED)
    c = by_check(claims, A, "flag")[0]
    assert c.confidence == "computed"
    assert cites(c, "02_financials/financial_model.xlsx", "ARR_Build!E13")
    assert not cites(c, CUST)
    assert any(g.what == "Not in the room: a revenue export by product type" for g in gaps)


def test_mrr_not_run_on_brasswick():  # D2
    assert outcomes(BRASS)[A] == "not_run"


# ---------------------------------------------------------------- kind b


def test_billings_clear_on_full_room():  # AM-R11, DC2
    claims = run(FULL)[0]
    assert not by_check(claims, B, "flag") and outcomes(FULL)[B] == "clear"
    info = by_check(claims, B, "info")[0]
    assert "21" in info.text


def test_billings_flags_brasswick_september():  # AM-R11
    c = by_check(run(BRASS)[0], B, "flag")
    assert len(c) == 1 and c[0].value.period == "2026-09"
    assert "$84,810" in c[0].text and "$82,650" in c[0].text and "$59,565" in c[0].text
    assert cites(c[0], "02_financials/financial_model.xlsx", "P&L!S7")


SEED_CANNOT = {B: "Not in the room: an invoices or billings export",
               C: "Not in the room: a revenue export that separates pass-through money",
               E: "Not in the room: an expense ledger",
               F: "Not in the room: an expense ledger",
               G: "Not in the room: a balance sheet with a deferred revenue line"}


@pytest.mark.parametrize("name", list(SEED_CANNOT))
def test_seed_drop_cannot_check_without_customer_file(name):  # AM-R12, R15, R20, R22, R25
    claims, gaps = run(SEED)
    assert outcomes(SEED)[name] == "cannot_check"
    assert not by_check(claims, name)
    assert any(g.what == SEED_CANNOT[name] and g.why == "file missing" for g in gaps)


# ---------------------------------------------------------------- kind c


def test_gross_vs_net_flags_brasswick_recoveries():  # AM-R13, AM-R14
    c = by_check(run(BRASS)[0], C, "flag")
    assert len(c) == 1
    assert "$264,214.66" in c[0].text and "$943,339.50" in c[0].text and "28.0%" in c[0].text
    assert cites(c[0], "03_metrics/revenue_by_customer.csv") and cites(c[0], "02_financials/expense_ledger.csv")


def test_gross_vs_net_fires_with_resale_on(resale_rooms):  # AM-R14
    _, on = resale_rooms
    c = by_check(chk(on)[0], C, "flag")
    assert len(c) == 1 and "2.7%" in c[0].text


def test_gross_vs_net_silent_with_resale_off(resale_rooms):  # AM-R14
    off, _ = resale_rooms
    claims, gaps = chk(off)
    assert not by_check(claims, C)
    assert mislabels.outcomes(off)[C] == "not_run"
    assert outcomes(FULL)[C] == "not_run"


def test_resale_fixture_writes_nothing_under_data_room(tmp_path_factory):  # AM-R30
    before = _tree_hash(FULL)
    base = tmp_path_factory.mktemp("again")
    _copy(base / "x")
    assert _tree_hash(FULL) == before


# ---------------------------------------------------------------- kind e


def test_contractor_lines_include_the_whole_vendor():  # AM-R16
    vendors = mislabels.contractor_vendors(BRASS)
    assert "Norrgate Software Studio" in vendors and "Fernway Contractor Payments" in vendors
    assert "Abbotsford Finance Partners" not in vendors  # the fractional controller, Brasswick DC2


def test_contractor_tie_clear_on_full_room():  # AM-R17
    assert outcomes(FULL)[E] == "clear"
    info = by_check(run(FULL)[0], E, "info")[0]
    assert "21" in info.text and "$0" in info.text


def test_contractors_cut_from_plan_flag_brasswick():  # AM-R18
    c = by_check(run(BRASS)[0], E, "flag")
    assert len(c) == 1
    assert "$34,708.97" in c[0].text and "$2,800" in c[0].text and "Professional services" in c[0].text
    assert cites(c[0], "02_financials/expense_ledger.csv")


def test_contractors_clear_with_gap_on_full_room():  # AM-R19
    assert any(g.what == "Not in the room: a list of contractors, or an org chart that shows them" for g in run(FULL)[1])


def test_contractors_found_without_flag_is_watch(monkeypatch):  # AM-R19
    # Brasswick with the plan keeping the contractors: the plan months of "Professional services"
    # carry the 2026-09 actual. Patched in memory, because saving the workbook drops cached values.
    original = mislabels.Room.pnl.fget

    def kept(self):
        p = original(self)
        row = p.rows["Professional services"][0]
        actual_col = next(col for m, s, col in p.months if m == "2026-09")
        for _m, s, col in p.months:
            if s == "Plan":
                p.values[(row, col)] = p.values[(row, actual_col)]
        return p

    monkeypatch.setattr(mislabels.Room, "pnl", property(kept))
    claims, _ = chk(BRASS)
    assert not by_check(claims, E, "flag")
    assert by_check(claims, E, "watch")


# ---------------------------------------------------------------- kind f


@pytest.mark.parametrize("room,bridge,cash", [(FULL, "-$0.24", "$1,800,000"), (BRASS, "$0.21", "$641,318")],
                         ids=["hollis_voice", "brasswick"])
def test_capitalised_clear_with_balance_sheet_gap(room, bridge, cash):  # AM-R21
    claims, gaps = run(room)
    assert outcomes(room)[F] == "clear"
    info = by_check(claims, F, "info")[0]
    assert bridge in info.text and cash in info.text
    assert any(g.what == "Not in the room: a balance sheet" for g in gaps)


# ---------------------------------------------------------------- kind g


def test_deferred_clear_on_full_room():  # AM-R23
    assert outcomes(FULL)[G] == "clear"
    info = by_check(run(FULL)[0], G, "info")[0]
    assert "$104,876" in info.text and "28 customers" in info.text and "$2,060,868" in info.text


def test_deferred_flags_ostley_on_brasswick():  # AM-R23
    c = by_check(run(BRASS)[0], G, "flag")
    assert len(c) == 1
    assert "$25,245" in c[0].text and "Ostley Freight" in c[0].text
    assert cites(c[0], "02_financials/financial_model.xlsx", "P&L!S7")


def test_bridge_deferred_is_info_when_pnl_flag_fired():  # AM-R24
    info = [c for c in by_check(run(BRASS)[0], G, "info") if "$33,885" in c.text]
    assert len(info) == 1 and "$8,640" in info[0].text


# ---------------------------------------------------------------- kind h


@pytest.mark.parametrize("room", [FULL, SEED], ids=["hollis_voice", "seed_drop"])
def test_one_time_fees_clear(room):  # AM-R26, DC1
    assert outcomes(room)[H] == "clear"
    info = by_check(run(room)[0], H, "info")[0]
    assert "$9,500" in info.text and "$15,000" in info.text


def test_one_time_fees_not_run_on_brasswick():  # AM-R26
    assert outcomes(BRASS)[H] == "not_run"


# ---------------------------------------------------------------- decoys, clean room, score


@pytest.mark.parametrize("room", [FULL, SEED, CLEAN], ids=lambda p: p.name)
def test_hollis_decoys_not_cited_by_any_flag(room):  # AM-R27
    for c in run(room)[0]:
        if c.severity not in ("flag", "watch"):
            continue
        for s in c.sources:
            assert ",C0193," not in s.quote
            assert ",Onboarding," not in s.quote
            assert not (s.uri == CUST and s.quote.startswith("2026-06,"))
            assert s.locator not in ("ARR_Build!A14", "ARR_Build!E14")
            assert "order_form_brookline_row" not in s.uri


def test_brasswick_decoys_not_cited_by_any_flag():  # AM-R27
    key = json.loads((ROOMS / "answer_key_brasswick.json").read_text(encoding="utf-8"))
    readout = Readout(room="brasswick", built_at="test", claims=run(BRASS)[0])
    line, _ = score(readout, key, "brasswick")
    assert line.decoys_flagged == 0


def test_clean_room_has_no_flag_or_watch():  # AM-R28
    assert not [c for c in run(CLEAN)[0] if c.severity in ("flag", "watch")]


@pytest.mark.parametrize("room,key", [(FULL, "answer_key.json"), (SEED, "answer_key.json"),
                                      (CLEAN, "answer_key_clean.json")], ids=lambda x: getattr(x, "name", x))
def test_eval_score_flags_no_hollis_decoy(room, key):  # AM-R29
    k = json.loads((ROOMS / key).read_text(encoding="utf-8"))
    line, _ = score(Readout(room=room.name, built_at="test", claims=run(room)[0]), k, room.name)
    assert line.decoys_flagged == 0


def test_eval_score_on_brasswick_finds_bf1_bf3_bf4():  # AM-R29
    key = json.loads((ROOMS / "answer_key_brasswick.json").read_text(encoding="utf-8"))
    _, lines = score(Readout(room="brasswick", built_at="test", claims=run(BRASS)[0]), key, "brasswick")
    found = {ln.split()[1] for ln in lines if ln.startswith("found")}
    assert {"BF1", "BF3", "BF4"} <= found


# ---------------------------------------------------------------- panels and the engine


def test_panels_carry_entry_checks_and_lines():  # entry 1.12 output shape, as Panels
    _, _, panels = mislabels.check(FULL, Ids("redflags"))
    by_name = {p.name: p for p in panels}
    entry = by_name["accounting mislabels: entry"]
    assert entry.rows[0][1] == "flag" and A in entry.rows[0][2]
    checks = by_name["accounting mislabels: checks"]
    assert [r[0] for r in checks.rows] == [A, B, C, E, F, G, H]
    assert dict((r[0], r[3]) for r in checks.rows)[C] == "not_run"
    lines = by_name["accounting mislabels: lines"]
    assert ["03_metrics/monthly_metrics.csv row 22", "mrr", "pilot"] == lines.rows[0][:3]
    assert any(r[0] == "P&L!W10" for r in lines.rows)


def test_seed_drop_entry_label():  # entry 1.12 scoring on the seed drop
    _, _, panels = mislabels.check(SEED, Ids("redflags"))
    assert {p.name: p for p in panels}["accounting mislabels: entry"].rows[0][1] == "flag"


def test_engine_result_holds_mislabels_claims_and_panels():  # joins redflag-engine
    from scans.redflags import scan
    result = scan.run(FULL)
    assert by_check(result.claims, A, "flag")
    assert any(p.name == "accounting mislabels: checks" for p in result.panels)
    assert len({c.id for c in result.claims}) == len(result.claims)
