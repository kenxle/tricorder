"""cohorts: the script half of FRAMEWORKS entry 1.3 (cohort retention rebuild).

Rooms in data_room/ are read in place and never written. Expected figures are
the answer key's PF3 (data_room/answer_key.json): the tab shows 102.6% over 57
customers; rebuilt, 87.3% NRR and 73.4% GRR over 68 customers.
"""

from __future__ import annotations

import shutil
from functools import lru_cache
from pathlib import Path

import pytest

from schema.check_quotes import check_claim
from scans.base import Ids
from scans.redflags import cohorts

DATA_ROOM = Path(__file__).resolve().parents[3] / "data_room"


@lru_cache(maxsize=None)
def run_all(name: str):
    return cohorts.check(DATA_ROOM / name, Ids("redflags"))


def run(name: str):
    claims, gaps, _panels = run_all(name)
    return claims, gaps


def by_label(claims, label):
    return [c for c in claims if c.value and c.value.label == label]


def test_seed_drop_states_the_tab_and_names_the_missing_customer_file():
    claims, gaps = run("hollis_voice_seed_drop")
    nrr = by_label(claims, "nrr_12m")
    assert len(nrr) == 1
    assert nrr[0].value.number == pytest.approx(102.56, abs=0.01)
    assert nrr[0].severity == "info"
    assert "57 customers" in nrr[0].text
    assert any("customer_mrr_by_month.csv" in g.what and g.piece == "retention" for g in gaps)
    assert not by_label(claims, "grr_12m")


def test_full_room_rebuilds_nrr_and_grr_and_flags_the_tab():
    claims, gaps = run("hollis_voice")
    tab = [c for c in by_label(claims, "nrr_12m") if "Cohorts tab states" in c.text]
    rebuilt = [c for c in by_label(claims, "nrr_12m") if c.text.startswith("Rebuilt")]
    grr = by_label(claims, "grr_12m")
    assert len(tab) == 1 and len(rebuilt) == 1 and len(grr) == 1
    assert rebuilt[0].value.number == pytest.approx(87.31, abs=0.01)
    assert rebuilt[0].severity == "flag"
    assert "68 customers" in rebuilt[0].text and "57" in rebuilt[0].text
    assert rebuilt[0].conflicts_with == [tab[0].id]
    assert grr[0].value.number == pytest.approx(73.39, abs=0.01)
    assert grr[0].severity == "watch"
    assert rebuilt[0].check.name == cohorts.CHECK
    assert not [g for g in gaps if "customer_mrr_by_month" in g.what]


def test_clean_room_agrees_and_does_not_flag():
    claims, _ = run("hollis_voice_clean")
    rebuilt = [c for c in by_label(claims, "nrr_12m") if c.text.startswith("Rebuilt")]
    assert rebuilt and rebuilt[0].severity == "info"
    assert rebuilt[0].conflicts_with == []
    assert not [c for c in claims if c.severity == "flag"]


@pytest.mark.parametrize("room", ["hollis_voice", "hollis_voice_seed_drop", "hollis_voice_clean"])
def test_every_claim_passes_the_quote_checker(room):
    claims, _ = run(room)
    assert claims
    for c in claims:
        _, problems = check_claim(c, DATA_ROOM / room)
        assert problems == [], (c.id, problems)


def test_no_tab_and_no_customer_file_is_a_gap_not_a_crash():
    claims, gaps = run("brasswick")
    assert claims == []
    assert any("Cohorts tab" in g.what for g in gaps)


def test_missing_model_file_is_a_gap(tmp_path):
    room = tmp_path / "room"
    (room / "03_metrics").mkdir(parents=True)
    shutil.copy(DATA_ROOM / "hollis_voice/03_metrics/customer_mrr_by_month.csv", room / "03_metrics")
    claims, gaps, _ = cohorts.check(room, Ids("redflags"))
    assert by_label(claims, "nrr_12m")[0].severity == "info"  # nothing to compare it with
    assert by_label(claims, "grr_12m")[0].severity == "watch"  # under the 80% line
    assert any("Cohorts tab" in g.what for g in gaps)


def test_listed_in_the_redflags_scan():
    from scans.redflags.scan import CHECKS

    assert "cohorts" in CHECKS


# A retention number opens the rows that produce it (fix round, both personas' blocker).


def test_rebuilt_lines_cite_the_script_the_header_and_the_rows_panel():
    claims, _gaps, panels = run_all("hollis_voice")
    rebuilt = [c for c in claims if c.text.startswith("Rebuilt")]
    assert len(rebuilt) == 2
    for c in rebuilt:
        assert "scans/redflags/cohorts.py" in c.text and f"'{cohorts.PANEL}'" in c.text
        csv_sources = [s for s in c.sources if s.uri == cohorts.CUSTOMERS]
        assert [s.locator for s in csv_sources] == ["row 1"]  # the header: the columns summed, not one sample row
    assert [p.name for p in panels] == [cohorts.PANEL]


def test_the_panel_lists_every_row_summed_and_adds_up_to_the_lines():
    import csv

    _claims, _gaps, panels = run_all("hollis_voice")
    panel = panels[0]
    lines = (DATA_ROOM / "hollis_voice" / cohorts.CUSTOMERS).read_text(encoding="utf-8").splitlines()
    header = lines[0].split(",")
    col = {name: i for i, name in enumerate(header)}
    total_start = total_end = customers = 0.0
    for row in panel.rows:
        cohort, n, start, end = row[0], int(row[1]), float(row[2]), float(row[3])
        start_rows = [int(x) for x in row[5].split(", ")]
        end_rows = [int(x) for x in row[6].split(", ")] if row[6] else []
        got = [next(csv.reader([lines[r - 1]])) for r in start_rows]
        assert len({g[col["customer_id"]] for g in got}) == n
        assert all(g[col["month"]] == cohort for g in got)
        assert sum(float(g[col["mrr"]]) for g in got) == start
        assert sum(float(next(csv.reader([lines[r - 1]]))[col["mrr"]]) for r in end_rows) == end
        total_start, total_end, customers = total_start + start, total_end + end, customers + n
    assert customers == 68
    assert round(100 * total_end / total_start, 1) == 87.3


def test_seed_drop_has_no_rows_panel():
    _claims, _gaps, panels = run_all("hollis_voice_seed_drop")
    assert panels == []


# Bad rows in the customer file (review wave 2a, O8): skipped and counted, never a crash -----------


def broken_room(tmp_path, edit):
    """A copy of the full room's customer file and model, with `edit(lines)` applied to the CSV lines."""
    room = tmp_path / "room"
    for rel in (cohorts.CUSTOMERS, cohorts.MODEL):
        (room / rel).parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(DATA_ROOM / "hollis_voice" / rel, room / rel)
    path = room / cohorts.CUSTOMERS
    lines = path.read_text(encoding="utf-8").splitlines()
    edit(lines)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return room


def blank(lines, n, column):
    """Blank one column of file row n (row 1 is the header)."""
    header = lines[0].split(",")
    cells = lines[n - 1].split(",")
    cells[header.index(column)] = ""
    lines[n - 1] = ",".join(cells)


def test_a_blank_month_and_a_text_mrr_are_skipped_and_named_in_a_gap(tmp_path):
    def edit(lines):
        blank(lines, 5, "month")
        header = lines[0].split(",")
        cells = lines[8].split(",")
        cells[header.index("mrr")] = "n/a"
        lines[8] = ",".join(cells)
    room = broken_room(tmp_path, edit)
    claims, gaps, panels = cohorts.check(room, Ids("redflags"))  # no ValueError
    assert by_label(claims, "nrr_12m") and by_label(claims, "grr_12m") and panels
    bad = [g for g in gaps if "customer file" in g.what and "skipped" in g.why]
    assert len(bad) == 1
    assert "2 of" in bad[0].why and "rows 5, 9" in bad[0].why
    rebuilt = [c for c in by_label(claims, "nrr_12m") if c.text.startswith("Rebuilt")][0]
    assert "skipped 2 rows" in rebuilt.text
    for c in claims:
        assert check_claim(c, room)[1] == []


def test_a_customer_file_with_no_readable_row_is_a_gap(tmp_path):
    def edit(lines):
        for n in range(2, len(lines) + 1):
            blank(lines, n, "month")
    room = broken_room(tmp_path, edit)
    claims, gaps, panels = cohorts.check(room, Ids("redflags"))
    assert [c for c in claims if c.text.startswith("Rebuilt")] == [] and panels == []
    assert any("customer file" in g.what and "skipped" in g.why for g in gaps)


def test_the_full_room_skips_nothing():
    _claims, gaps, _ = run_all("hollis_voice")
    assert not [g for g in gaps if "skipped" in g.why]


# A tab far from the rebuild is never "matches" (review wave 2b, S9) -------------------------------


def tab_room(tmp_path, customers, nrr):
    """A copy of the full room's customer file and model, with the Cohorts tab's "All cohorts" row set."""
    from openpyxl import load_workbook

    room = broken_room(tmp_path, lambda lines: None)
    path = room / cohorts.MODEL
    wb = load_workbook(path)
    ws = wb[cohorts.TAB]
    hdr = next(r for r in ws.iter_rows() if isinstance(r[0].value, str) and r[0].value.startswith("Cohort ("))
    names = [c.value for c in hdr]
    total = next(r for r in ws.iter_rows() if r[0].value == "All cohorts")
    total[names.index("Customers")].value = customers
    total[names.index("Net revenue retention")].value = nrr
    wb.save(path)
    return room


def rebuilt_line(room):
    claims, _, _ = cohorts.check(room, Ids("redflags"))
    for c in claims:
        assert check_claim(c, room)[1] == []
    return next(c for c in by_label(claims, "nrr_12m") if c.text.startswith("Rebuilt"))


def test_a_tab_far_below_the_rebuild_is_not_a_match(tmp_path):
    line = rebuilt_line(tab_room(tmp_path, 68, 0.80))
    assert "matches" not in line.text
    assert "7.3 points higher than the Cohorts tab's 80.0% over 68 customers" in line.text
    assert line.severity == "watch" and line.conflicts_with == []


def test_a_tab_within_tolerance_with_the_same_count_matches(tmp_path):
    line = rebuilt_line(tab_room(tmp_path, 68, 0.875))
    assert "which matches the Cohorts tab's 87.5% over 68 customers" in line.text
    assert line.severity == "info"


def test_a_tab_within_tolerance_with_more_customers_says_so(tmp_path):
    line = rebuilt_line(tab_room(tmp_path, 70, 0.875))
    assert "matches" not in line.text and "the tab counts 2 customers more" in line.text
    assert line.severity == "info"
