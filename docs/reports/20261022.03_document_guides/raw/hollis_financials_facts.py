"""Facts about the Hollis Voice room that the financials guides quote.

Run with the data room's own interpreter so openpyxl is available:
  data_room/.venv/bin/python docs/reports/20261022.03_document_guides/raw/hollis_financials_facts.py
Reads data_room/hollis_voice (the main room) and data_room/answer_key.json. Prints one JSON object.
Every number in the guides' "In the Hollis Voice room" sections comes from this output
(saved as raw/hollis_financials_facts_output.json).
"""
import csv, json, pathlib, statistics
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[4] / "data_room"
ROOM = ROOT / "hollis_voice"
out = {}

wb = openpyxl.load_workbook(ROOM / "02_financials" / "financial_model.xlsx", data_only=False)
wbv = openpyxl.load_workbook(ROOM / "02_financials" / "financial_model.xlsx", data_only=True)

out["model_tabs"] = wb.sheetnames
out["model_tab_count"] = len(wb.sheetnames)
out["named_cells"] = len(wb.defined_names)
formulas = {}
for ws in wb.worksheets:
    n = 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n += 1
    formulas[ws.title] = n
out["formulas_by_tab"] = formulas
out["formulas_total"] = sum(formulas.values())

text_all = []
for ws in wbv.worksheets:
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if isinstance(v, str):
                text_all.append(v.lower())
blob = " ".join(text_all)
out["words_found_in_model"] = {w: (w in blob) for w in ["scenario", "downside", "sensitivity", "balance sheet", "cash flow statement", "deferred revenue", "working capital"]}
out["tabs_named_balance_sheet_or_cash_flow"] = [t for t in wb.sheetnames if "balance" in t.lower() or "cash_flow" in t.lower() or "cash flow" in t.lower()]

# Assumptions tab: how many inputs carry a note
a = wbv["Assumptions"]
inputs = [r for r in a.iter_rows(min_row=1, values_only=True) if r[1] is not None and isinstance(r[1], (int, float, str)) and r[0] not in (None, "Input", "Assumptions")]
out["assumption_rows"] = len(inputs)
out["assumption_rows_with_note"] = sum(1 for r in inputs if r[2])

# P&L: actual and plan months
pl = wbv["P&L"]
hdr = [c.value for c in pl[4]]
status = [c.value for c in pl[5]]
months = [(h, s) for h, s in zip(hdr, status) if isinstance(h, str) and h[:2] == "20"]
out["pnl_actual_months"] = sum(1 for h, s in months if s == "Actual")
out["pnl_plan_months"] = sum(1 for h, s in months if s == "Plan")
out["pnl_first_month"] = months[0][0]
out["pnl_last_month"] = months[-1][0]

# Cash_Runway: months in the cash plan, payroll against revenue
cr = wbv["Cash_Runway"]
rows = [r for r in cr.iter_rows(min_row=9, values_only=True) if isinstance(r[1], int)]
out["cash_plan_months"] = len(rows)
rev = [r[2] for r in rows]
payroll = [r[4] for r in rows]
other = [r[5] for r in rows]
mkt = [r[6] for r in rows]
out["cash_plan_first_last_month"] = [rows[0][0], rows[-1][0]]
out["revenue_month1"] = round(rev[0], 2)
out["revenue_month60"] = round(rev[-1], 2)
out["revenue_growth_multiple_month1_to_last"] = round(rev[-1] / rev[0], 2)
out["payroll_distinct_values_from_plan_month_4"] = len(set(round(p, 2) for p in payroll[3:]))
out["payroll_month4"] = round(payroll[3], 2)
out["payroll_last"] = round(payroll[-1], 2)
out["other_costs_distinct_values"] = len(set(round(o, 2) for o in other))
out["marketing_programs_monthly_after_raise"] = sorted(set(round(m, 2) for m in mkt if m))
out["runway_months_in_model"] = cr["B3"].value
out["avg_net_burn_12m_in_model"] = round(cr["B4"].value, 2)

# Hiring plan versus the model
hp = list(csv.DictReader(open(ROOM / "05_team" / "hiring_plan.csv")))
out["hiring_plan_rows"] = len(hp)
out["hires_counted_in_model"] = wbv["Hiring_Plan"]["B13"].value if wbv["Hiring_Plan"]["B13"].value else None
by_dept = {}
for r in hp:
    by_dept[r["department"]] = by_dept.get(r["department"], 0) + 1
out["hires_by_department"] = by_dept
LOAD = 1.28
out["plan_annual_loaded_cost_by_department_usd"] = {
    d: round(sum(float(r["base_salary_usd"]) for r in hp if r["department"] == d) * LOAD) for d in by_dept}

# Ledger against the P&L
ledger = list(csv.DictReader(open(ROOM / "02_financials" / "expense_ledger.csv")))
out["ledger_lines"] = len(ledger)
led_by_month = {}
for r in ledger:
    led_by_month[r["month"]] = led_by_month.get(r["month"], 0) + float(r["amount_usd"])
# P&L cost totals by month: rows 15 (total cost of revenue) and 26 (total opex)
hdrs = {h: i for i, h in enumerate(hdr) if isinstance(h, str) and h[:2] == "20"}
gaps = {}
for m, i in hdrs.items():
    if m in led_by_month:
        cogs = pl.cell(15, i + 1).value
        opex = pl.cell(26, i + 1).value
        gaps[m] = round(led_by_month[m] - (cogs + opex), 2)
out["ledger_vs_pnl_months_compared"] = len(gaps)
out["ledger_vs_pnl_max_abs_gap_usd"] = max(abs(g) for g in gaps.values()) if gaps else None

# P&L subscription revenue against the customer file
cust = list(csv.DictReader(open(ROOM / "03_metrics" / "customer_mrr_by_month.csv")))
sub_by_month = {}
for r in cust:
    if r["product_type"] == "Subscription":
        sub_by_month[r["month"]] = sub_by_month.get(r["month"], 0) + float(r["mrr"])
sub_gaps = []
for m, i in hdrs.items():
    if m in sub_by_month:
        sub_gaps.append(abs(pl.cell(7, i + 1).value - sub_by_month[m]))
out["pnl_subscription_vs_customer_file_months"] = len(sub_gaps)
out["pnl_subscription_vs_customer_file_max_gap_usd"] = round(max(sub_gaps), 2) if sub_gaps else None

# Gross margin as shown and as rebuilt for 2026-09 from the ledger
sep = "2026-09"
cogs_cats_true = {"telephony", "third_party_licences", "cloud_hosting", "model_inference", "support_staff"}
led_sep = [r for r in ledger if r["month"] == sep]
true_cogs = sum(float(r["amount_usd"]) for r in led_sep if r["category"] in cogs_cats_true)
rec_rev = None
mm = {r["month"]: r for r in csv.DictReader(open(ROOM / "03_metrics" / "monthly_metrics.csv"))}
i = hdrs[sep]
rec_rev = pl.cell(10, i + 1).value
out["gross_margin_shown_sep_2026"] = round(pl.cell(17, i + 1).value, 4)
out["gross_margin_rebuilt_from_ledger_sep_2026"] = round((rec_rev - true_cogs) / rec_rev, 4)

# Cash bridge
cb = wbv["Cash_Bridge"]
out["cash_bridge_difference_usd"] = round([r for r in cb.iter_rows(values_only=True) if r[0] == "Difference"][0][1], 2)

# Use of funds on slide 13 (text read with pdftotext, entered here from the slide)
uof = {"Engineering": 3600000, "Sales and marketing": 3150000, "Customer success and support": 1350000, "General and administrative": 900000}
raise_usd = 9000000
out["use_of_funds_sum_usd"] = sum(uof.values())
out["use_of_funds_sum_equals_raise"] = sum(uof.values()) == raise_usd
out["use_of_funds_shares_percent"] = {k: round(100 * v / raise_usd, 2) for k, v in uof.items()}
out["use_of_funds_categories_missing_from_pnl_lines"] = ["Customer success and support"]

# Tax years a company founded in 2024 would have filed by the data date (calendar-year C corporation)
founded, data_year = 2024, 2026
out["tax_years_ended_before_data_date"] = list(range(founded, data_year))
out["tax_returns_due_by_data_date_count"] = len(range(founded, data_year))


# Extra model checks: plan columns hard-coded or formula, totals foot, growth in the plan against recent actual growth
pl_f = wb["P&L"]
plan_cols = [i + 1 for (h, st), i in zip(months, [k for k, hh in enumerate(hdr) if isinstance(hh, str) and hh[:2] == "20"]) if st == "Plan"]
const_in_plan = formula_in_plan = 0
for col in plan_cols:
    for r in range(7, 32):
        v = pl_f.cell(r, col).value
        if isinstance(v, str) and v.startswith("="):
            formula_in_plan += 1
        elif isinstance(v, (int, float)):
            const_in_plan += 1
out["pnl_plan_cells_formula"] = formula_in_plan
out["pnl_plan_cells_hardcoded_number"] = const_in_plan

foot_bad = 0
for m, i in hdrs.items():
    c = i + 1
    cogs = pl.cell(13, c).value + pl.cell(14, c).value
    if abs(cogs - pl.cell(15, c).value) > 0.01:
        foot_bad += 1
    opex = sum(pl.cell(r, c).value for r in (19, 20, 21, 22, 23, 24, 25))
    if abs(opex - pl.cell(26, c).value) > 0.01:
        foot_bad += 1
out["pnl_months_checked_for_footing"] = len(hdrs)
out["pnl_subtotals_that_do_not_foot"] = foot_bad

mrr_rows = list(csv.DictReader(open(ROOM / "03_metrics" / "monthly_metrics.csv")))
last = [float(r["mrr"]) for r in mrr_rows[-7:]]
cmgr = (last[-1] / last[0]) ** (1 / 6) - 1
out["actual_mrr_growth_per_month_last_6_months_percent"] = round(100 * cmgr, 2)
out["plan_growth_per_month_percent"] = 3.5

# Files the room holds of the six kinds
files = json.load(open(ROOT / "answer_key.json"))["full_room_files"]
out["room_file_count"] = len(files)
out["room_has_bank_statements"] = any("bank" in f.lower() for f in files)
out["room_has_tax_returns"] = any("tax" in f.lower() for f in files)
out["room_has_audited_or_reviewed"] = any("audit" in f.lower() or "review" in f.lower() for f in files)
out["room_has_balance_sheet_file"] = any("balance" in f.lower() for f in files)

print(json.dumps(out, indent=1))
