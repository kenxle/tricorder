#!/usr/bin/env python3
"""Spike: deck revenue against the books, and partial history.

Standalone. Not the production module. It reads a data room the way the
production check will (deck by page, model P&L by cell, revenue rows by row),
matches every deck revenue figure to the same period in the books, and checks
whether a deck revenue chart shows only part of the history.

Usage:
    python3 deck_vs_books_spike.py <room_dir> [--deck PATH] [--hide FILE ...]
                                    [--out PATH]

--deck   reads a different deck (for example a test deck made by
         make_test_decks.py) in place of <room>/01_overview/pitch_deck.pdf.
--hide   pretends a room file is missing (relative path), to show the
         failure cases without copying or changing the room.
--out    writes the check output (claims, gaps, matches, unmatched) as JSON.

Needs pdfplumber, pypdf, openpyxl and Pillow (data_room/requirements.txt has
pdfplumber and openpyxl; pypdf and Pillow come with them here).
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.utils import get_column_letter
from PIL import Image
from pypdf import PdfReader

SCAN = "redflags"
CHECK_REVENUE = "deck_revenue_vs_pnl"
CHECK_HISTORY = "deck_history_vs_metrics"  # the name answer_key_brasswick.json gives BF2
MATCH_TOLERANCE = 0.01      # FRAMEWORKS.md "Settings", match_tolerance
CHART_TOLERANCE = 0.02      # spike setting: bar height read from pixels, as a share of the tallest bar
DECK = "01_overview/pitch_deck.pdf"
CLAIMS_JSON = "01_overview/claims.json"
MODEL = "02_financials/financial_model.xlsx"
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
TODAY = date.today().isoformat()


# ----------------------------------------------------------------------------- helpers
def ym(year: int, month: int) -> str:
    return f"{year:04d}-{month:02d}"


def month_name_to_ym(text: str) -> str:
    name, year = text.split()
    return ym(int(year), MONTHS.index(name) + 1)


def month_range(start: str, end: str) -> list[str]:
    y, m = map(int, start.split("-"))
    out = []
    while ym(y, m) <= end:
        out.append(ym(y, m))
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def quarter_of(month: str) -> str:
    y, m = map(int, month.split("-"))
    return f"{y}-Q{(m - 1) // 3 + 1}"


def parse_money(text: str) -> tuple[float, float]:
    """'$2.5M' -> (2500000, 50000): the value and half a unit of its last printed digit."""
    t = text.strip().replace(",", "")
    mult = {"K": 1e3, "M": 1e6, "B": 1e9}
    m = re.fullmatch(r"\$?([\d.]+)\s*([KMB])?", t)
    if not m:
        raise ValueError(text)
    digits, suffix = m.group(1), m.group(2)
    scale = mult.get(suffix, 1)
    decimals = len(digits.split(".")[1]) if "." in digits else 0
    half_step = 0.5 * scale * 10 ** (-decimals) if suffix else 0.5
    return float(digits) * scale, half_step


def money(v: float) -> str:
    return f"${v:,.0f}"


def quote_of(value) -> str:
    """The text a cell reader is assumed to return for a number (assumption A4)."""
    return str(value)


# ----------------------------------------------------------------------------- room readers
class Room:
    def __init__(self, root: Path, deck: Path | None, hidden: set[str]):
        self.root = root
        self.deck_path = deck or root / DECK
        self.deck_uri = DECK if deck is None else str(deck.resolve())
        self.hidden = hidden

    def has(self, rel: str) -> bool:
        return rel not in self.hidden and (self.root / rel).exists()

    def deck_ok(self) -> bool:
        return self.deck_path.exists() and (str(self.deck_path) not in self.hidden and DECK not in self.hidden
                                            or self.deck_uri != DECK)


def read_deck(path: Path) -> list[dict]:
    pages = []
    reader = PdfReader(str(path))
    with pdfplumber.open(str(path)) as pdf:
        for i, p in enumerate(pdf.pages, 1):
            imgs = []
            for im in reader.pages[i - 1].images:
                try:
                    imgs.append(Image.open(io.BytesIO(im.data)).convert("RGB"))
                except Exception:  # an image we cannot decode counts as unreadable
                    pass
            pages.append({"slide": i, "text": p.extract_text() or "", "images": imgs})
    return pages


def read_pnl(path: Path) -> dict:
    """Rows of the model's P&L tab, keyed by line label, with their cell addresses."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["P&L"]
    header_row = next(r for r in range(1, 15) if ws.cell(r, 1).value == "Line")
    months, kinds = {}, {}
    for col in range(3, ws.max_column + 1):
        v = ws.cell(header_row, col).value
        if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}", v):
            months[v] = col
            kinds[v] = ws.cell(header_row + 1, col).value
    lines = {}
    for r in range(header_row + 2, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label and ws.cell(r, 2).value == "Revenue":
            lines[label] = {m: {"value": ws.cell(r, c).value,
                                "ref": f"P&L!{get_column_letter(c)}{r}"} for m, c in months.items()}
            lines[label]["_row"] = r
    named = {}
    for name, d in wb.defined_names.items():
        for sheet, ref in d.destinations:
            named[name] = wb[sheet][ref.replace("$", "")].value
    return {"months": months, "kinds": kinds, "lines": lines, "named": named}


def revenue_rows_by_month(room: Room) -> tuple[str, dict] | None:
    """Sum of the customer-level revenue rows per month, for a second read of the books."""
    for rel, cols in (("03_metrics/customer_mrr_by_month.csv", ["mrr"]),
                      ("03_metrics/revenue_by_customer.csv", ["subscription_revenue", "recovered_from_carriers"])):
        if room.has(rel):
            sums: dict[str, float] = {}
            with open(room.root / rel, newline="") as f:
                for row in csv.DictReader(f):
                    sums[row["month"]] = sums.get(row["month"], 0.0) + sum(float(row[c]) for c in cols)
            return rel, sums
    return None


# ----------------------------------------------------------------------------- chart pixels
def read_bars(img: Image.Image) -> list[float] | None:
    """Bar heights in pixels, left to right, for a single-colour bar chart. None if unreadable."""
    w, h = img.size
    px = img.load()
    counts = Counter(px[x, y] for x in range(0, w, 2) for y in range(0, h, 2))
    colours = [c for c, n in counts.most_common(6) if c != (255, 255, 255) and max(c) - min(c) > 40]
    if not colours:
        return None
    bar = colours[0]

    def is_bar(x, y):
        r, g, b = px[x, y]
        return abs(r - bar[0]) + abs(g - bar[1]) + abs(b - bar[2]) < 30

    base = max((y for y in range(h) for x in range(0, w, 4) if is_bar(x, y)), default=None)
    if base is None:
        return None
    row = base - 1
    runs, start = [], None
    for x in range(w):
        if is_bar(x, row) and start is None:
            start = x
        elif not is_bar(x, row) and start is not None:
            if x - start >= 4:
                runs.append((start, x - 1))
            start = None
    heights = []
    for a, b in runs:
        cx = (a + b) // 2
        top = base
        while top > 0 and is_bar(cx, top - 1):
            top -= 1
        heights.append(base - top + 1)
    return heights or None


# ----------------------------------------------------------------------------- claim builders
class Out:
    def __init__(self):
        self.claims, self.gaps, self.matches, self.unmatched = [], [], [], []
        self.n = 0

    def claim(self, text, value, severity, check, sources):
        self.n += 1
        cid = f"redflags-dvb-{self.n:03d}"
        self.claims.append({
            "id": cid, "scan": SCAN, "text": text, "value": value, "severity": severity,
            "check": {"kind": "script", "name": check, "model": None},
            "confidence": "computed", "sources": sources,
            "conflicts_with": [], "corroborated_by": [],
        })
        return cid

    def gap(self, what, why, would_need):
        self.gaps.append({"what": what, "why": why, "would_need": would_need})


def src_file(uri, locator, quote):
    return {"kind": "file", "uri": uri, "locator": locator, "start_s": None, "end_s": None,
            "quote": quote, "retrieved_at": TODAY}


def value(number, unit, label, period):
    return {"number": number, "unit": unit, "label": label, "period": period}


# ----------------------------------------------------------------------------- the check
def find_deck_figures(pages: list[dict], room: Room) -> tuple[list[dict], str | None]:
    """Every revenue figure the deck shows: claims-page rows, charts, growth sentences."""
    figures = []
    as_of = None
    m = re.search(r"Data as of (\d{4}-\d{2})-\d{2}", pages[0]["text"])
    if m:
        as_of = m.group(1)
    # Claims page: the last slide whose title starts with "Appendix".
    claims_slide = next((p for p in reversed(pages) if p["text"].startswith("Appendix")), None)
    rows = []
    if claims_slide:
        for line in claims_slide["text"].splitlines():
            mm = re.match(r"(.+?)\s+(\$[\d,.]+[KMB]?)\s+(\d+)$", line)
            if mm and re.search(r"\b(ARR|MRR|[Rr]evenue)\b", mm.group(1)):
                rows.append({"label": mm.group(1).strip(), "value": mm.group(2), "slide": int(mm.group(3)),
                             "row_text": line, "claims_slide": claims_slide["slide"]})
    elif room.has(CLAIMS_JSON):  # FRAMEWORKS 1.1: claims.json is the fallback
        data = json.loads((room.root / CLAIMS_JSON).read_text())
        for r in data["claims"]:
            if re.search(r"ARR|MRR|[Rr]evenue", r["label"]) and r["value"].startswith("$"):
                rows.append({"label": r["label"], "value": r["value"], "slide": r["slide"],
                             "row_text": None, "claims_slide": None})
    for r in rows:
        number, _ = parse_money(r["value"])
        short = None  # the figure as printed on its content slide, for example "$2.5M"
        content = pages[r["slide"] - 1]["text"]
        for tok in re.findall(r"\$[\d,.]+[KMB]?", content):
            v, half = parse_money(tok)
            if abs(v - number) <= max(half, 0.5):
                short = tok
                break
        figures.append({"kind": "stated", **r, "number": number, "short": short})
    # Charts: a caption naming a revenue series and a window.
    for p in pages:
        for line in p["text"].splitlines():
            mm = re.match(r"(Monthly|Quarterly) (recurring )?revenue, (\w+ \d{4}) to (\w+ \d{4})", line)
            if mm:
                figures.append({"kind": "chart", "slide": p["slide"], "caption": line.strip(),
                                "grain": mm.group(1).lower(), "recurring": bool(mm.group(2)),
                                "start": month_name_to_ym(mm.group(3)), "end": month_name_to_ym(mm.group(4)),
                                "image": p["images"][0] if p["images"] else None})
    # Growth sentences.
    for p in pages:
        for line in p["text"].splitlines():
            mm = re.search(r"Revenue grows ([\d.]+)% a month in the plan\.", line)
            if mm:
                figures.append({"kind": "plan_growth", "slide": p["slide"], "quote": mm.group(0),
                                "pct": float(mm.group(1))})
            mm = re.search(r"Revenue grew every month in (\d{4})\.", line)
            if mm:
                figures.append({"kind": "grew_every_month", "slide": p["slide"], "quote": mm.group(0),
                                "year": int(mm.group(1))})
    return figures, as_of


def run(room: Room) -> dict:
    out = Out()
    if not room.deck_path.exists() or room.deck_uri == DECK and DECK in room.hidden:
        out.gap("Not in the room: the pitch deck (01_overview/pitch_deck.pdf)", "file missing",
                "The pitch deck")
        return finish(out, "cannot_check", "No deck, so no deck figure to match.")
    pages = read_deck(room.deck_path)
    figures, as_of = find_deck_figures(pages, room)

    if not room.has(MODEL):
        out.gap("Not in the room: the monthly P&L (02_financials/financial_model.xlsx, P&L tab)",
                "file missing", "The monthly income statement for every month since the first revenue")
        rr = revenue_rows_by_month(room)
        for f in figures:
            out.unmatched.append({"slide": f["slide"], "kind": f["kind"],
                                  "why": "no P&L to match against" + (f"; revenue rows in {rr[0]} could stand in" if rr else "")})
        return finish(out, "cannot_check", "The P&L is missing, so no deck figure can be matched to a period.")

    pnl = read_pnl(room.root / MODEL)
    lines = pnl["lines"]
    actual = [m for m, k in pnl["kinds"].items() if k == "Actual"]
    total = lines.get("Total revenue")
    rec_label = next((l for l in lines if l.startswith("Recurring revenue")), "Subscription revenue")
    recurring = lines[rec_label]
    rr = revenue_rows_by_month(room)

    def books_cell(line, month):
        return line[month]["value"], line[month]["ref"]

    first_rev = next(m for m in actual if (total[m]["value"] or 0) > 0)

    for f in figures:
        if f["kind"] == "stated":
            deck_sources = []
            if f["short"]:
                deck_sources.append(src_file(room.deck_uri, f"slide {f['slide']}", f["short"]))
            if f["row_text"]:
                deck_sources.append(src_file(room.deck_uri, f"slide {f['claims_slide']}", f["row_text"]))
            label = f["label"]
            mm = re.search(r"(\w+) (\d{4})$", label)
            if label == "ARR":
                period = as_of
                b, ref = books_cell(recurring, period)
                books = 12 * b
                basis = f"12 x '{rec_label}' for {period}"
                vlabel, unit = "arr", "USD_per_year"
                bsrc = [src_file(MODEL, ref, quote_of(b))]
            elif label.startswith("MRR") and mm:
                period = month_name_to_ym(f"{mm.group(1)} {mm.group(2)}")
                line = lines.get("Subscription revenue")
                b, ref = books_cell(line, period)
                books, basis = b, f"'Subscription revenue' for {period}"
                vlabel, unit = "mrr", "USD_per_month"
                bsrc = [src_file(MODEL, ref, quote_of(b))]
            elif "last 12 months" in label:
                period = f"{actual[-12]} to {actual[-1]}" if as_of == actual[-1] else None
                months12 = actual[-12:]
                books = sum(total[m]["value"] for m in months12)
                basis = f"sum of 'Total revenue', {months12[0]} to {months12[-1]}"
                vlabel, unit = "other", "USD"
                bsrc = [src_file(MODEL, total[m]["ref"], quote_of(total[m]["value"])) for m in months12]
            else:
                out.unmatched.append({"slide": f["slide"], "kind": "stated", "figure": f["value"],
                                      "why": f"no rule maps the label '{label}' to a P&L line"})
                continue
            half = parse_money(f["short"])[1] if f["short"] and not f["row_text"] else 0
            gap_rel = abs(f["number"] - books) / abs(books)
            ok = gap_rel <= MATCH_TOLERANCE or abs(f["number"] - books) <= half
            out.matches.append({"slide": f["slide"], "figure": f["value"], "period": period, "books": round(books, 2),
                                "basis": basis, "gap_rel_pct": round(100 * gap_rel, 1), "within_tolerance": ok})
            text = (f"The deck's {label} of {money(f['number'])} (slide {f['slide']}) "
                    + ("matches" if ok else "differs from")
                    + f" the P&L's {money(books)} ({basis}), a gap of {100 * gap_rel:.1f}%.")
            out.claim(text, value(f["number"], unit, vlabel, period), "info" if ok else "flag",
                      CHECK_REVENUE, deck_sources + bsrc)

        elif f["kind"] == "chart":
            line = recurring if f["recurring"] else total
            lname = rec_label if f["recurring"] else "Total revenue"
            window = month_range(f["start"], f["end"])
            cap_src = src_file(room.deck_uri, f"slide {f['slide']}", f["caption"])
            bars = read_bars(f["image"]) if f["image"] is not None else None
            # Months the chart claims that the books do not hold: a gap for the ask list.
            missing = [m for m in window if m not in actual]
            if missing:
                out.gap(f"Not in the room: monthly revenue for {missing[0]} to {missing[-1]}, shown on slide {f['slide']}",
                        "not stated in the file", f"The monthly P&L for {missing[0]} to {missing[-1]}")
            if bars is None:
                out.gap(f"Not in the room: a readable chart on slide {f['slide']}", "chart image could not be read",
                        "The numbers behind the chart, or a readable copy of the slide")
                expected_bars = None
            else:
                expected_bars = len(window) if f["grain"] == "monthly" else len({quarter_of(m) for m in window})
            # Bar shape against the books: each bar as a share of the tallest, scale free.
            if bars is not None and not missing:
                if f["grain"] == "monthly":
                    series = [(m, line[m]["value"]) for m in window]
                else:
                    q: dict[str, float] = {}
                    for m in window:
                        qk = quarter_of(m)
                        q[qk] = q.get(qk, 0) + line[m]["value"]
                    series = list(q.items())
                shape_ok = len(bars) == len(series)
                worst = None
                if shape_ok:
                    top_bar, top_book = max(bars), max(v for _, v in series)
                    diffs = [abs(b / top_bar - v / top_book) for b, (_, v) in zip(bars, series)]
                    worst = max(diffs)
                    shape_ok = worst <= CHART_TOLERANCE
                out.matches.append({"slide": f["slide"], "figure": f"chart, {len(bars)} bars", "period": f"{window[0]} to {window[-1]}",
                                    "books": f"'{lname}', {len(series)} {'months' if f['grain']=='monthly' else 'quarters'}",
                                    "basis": "bar height / tallest bar against month / largest month",
                                    "gap_rel_pct": None if worst is None else round(100 * worst, 1),
                                    "within_tolerance": shape_ok})
                peak_m, peak_v = max(series, key=lambda s: s[1])
                ref = line[peak_m]["ref"] if f["grain"] == "monthly" else line[window[-1]]["ref"]
                out.claim(
                    f"The deck's revenue chart on slide {f['slide']} has {len(bars)} bars for {window[0]} to {window[-1]}; "
                    f"their shape {'matches' if shape_ok else 'does not match'} '{lname}' in the P&L"
                    + (f" (largest difference {100 * worst:.1f} points of the tallest bar)." if worst is not None else "."),
                    None, "info" if shape_ok else "flag", CHECK_REVENUE,
                    [cap_src, src_file(MODEL, line[window[0]]["ref"], quote_of(line[window[0]]["value"])),
                     src_file(MODEL, ref, quote_of(line[peak_m]["value"]) if f["grain"] == "monthly" else quote_of(line[window[-1]]["value"]))])
            # Partial history.
            left_out = [m for m in actual if m < window[0] and m >= first_rev]
            skipped = (expected_bars - len(bars)) if (bars is not None and expected_bars and f["grain"] == "monthly") else 0
            quarters = f["grain"] == "quarterly"
            if left_out or skipped > 0 or quarters:
                held = describe_months(left_out, line) if left_out else None
                parts = []
                if left_out:
                    parts.append(f"starts at {window[0]}, {len(left_out)} months after the first revenue month in the books "
                                 f"({first_rev}); the months left out, {left_out[0]} to {left_out[-1]}, {held['text']}")
                if skipped > 0:
                    parts.append(f"shows {len(bars)} bars for a {expected_bars}-month window, so {skipped} months are skipped")
                if quarters:
                    dips = [m for i, m in enumerate(window[1:], 1) if line[m]["value"] < line[window[i - 1]]["value"]]
                    parts.append("shows quarters where the books hold months"
                                 + (f"; the quarters hide month-on-month falls in {', '.join(dips)}" if dips
                                    else "; no month inside a quarter fell"))
                srcs = [cap_src]
                if left_out:
                    srcs.append(src_file(MODEL, line[first_rev]["ref"], quote_of(line[first_rev]["value"])))
                    pk = held["peak_month"]
                    if pk != first_rev:
                        srcs.append(src_file(MODEL, line[pk]["ref"], quote_of(line[pk]["value"])))
                    srcs.append(src_file(MODEL, line[window[0]]["ref"], quote_of(line[window[0]]["value"])))
                out.claim(f"Shows part of the history: the revenue chart on slide {f['slide']} " + "; it ".join(parts) + ".",
                          value(len(left_out) + max(skipped, 0), "months", "other", f"{left_out[0]} to {left_out[-1]}" if left_out else None),
                          "flag", CHECK_HISTORY, srcs)
                out.matches.append({"slide": f["slide"], "figure": "chart window", "period": f"{window[0]} to {window[-1]}",
                                    "books": f"revenue from {first_rev}", "basis": "first chart month against first revenue month",
                                    "gap_rel_pct": None, "within_tolerance": False,
                                    "left_out": left_out, "left_out_hold": held})
            else:
                out.matches.append({"slide": f["slide"], "figure": "chart window", "period": f"{window[0]} to {window[-1]}",
                                    "books": f"revenue from {first_rev}", "basis": "first chart month against first revenue month",
                                    "gap_rel_pct": None, "within_tolerance": True})
                out.claim(f"The revenue chart on slide {f['slide']} starts at {window[0]}, the first revenue month in the books, "
                          f"and runs to {window[-1]} with no month left out.",
                          value(0, "months", "other", f"{window[0]} to {window[-1]}"), "info", CHECK_HISTORY,
                          [cap_src, src_file(MODEL, line[first_rev]["ref"], quote_of(line[first_rev]["value"]))])

        elif f["kind"] == "plan_growth":
            plan = [m for m, k in pnl["kinds"].items() if k == "Plan"]
            seq = [actual[-1]] + plan
            if any(not isinstance(recurring[m]["value"], (int, float)) for m in seq):
                out.unmatched.append({"slide": f["slide"], "kind": "plan_growth", "figure": f"{f['pct']}% a month",
                                      "why": f"the plan columns of '{rec_label}' hold no cached values"})
                continue
            rates = [recurring[b]["value"] / recurring[a]["value"] - 1 for a, b in zip(seq, seq[1:])]
            lo, hi = 100 * min(rates), 100 * max(rates)
            ok = all(abs(100 * r - f["pct"]) <= f["pct"] * MATCH_TOLERANCE for r in rates)
            out.matches.append({"slide": f["slide"], "figure": f"{f['pct']}% a month (plan)", "period": f"{plan[0]} to {plan[-1]}",
                                "books": f"plan columns {lo:.2f}% to {hi:.2f}%", "basis": f"month-on-month growth of '{rec_label}', plan columns",
                                "gap_rel_pct": None, "within_tolerance": ok})
            out.claim(f"The deck's plan growth of {f['pct']}% a month (slide {f['slide']}) "
                      f"{'matches' if ok else 'differs from'} the model's plan columns, which grow {lo:.2f}% to {hi:.2f}% a month. "
                      f"It is a plan, not a period the books hold as actuals.",
                      value(f["pct"], "percent", "other", f"{plan[0]} to {plan[-1]}"), "info", CHECK_REVENUE,
                      [src_file(room.deck_uri, f"slide {f['slide']}", f["quote"]),
                       src_file(MODEL, recurring[plan[0]]["ref"], quote_of(recurring[plan[0]]["value"]))])

        elif f["kind"] == "grew_every_month":
            yr = [m for m in actual if m.startswith(str(f["year"]))]
            falls = [m for a, m in zip(yr, yr[1:]) if total[m]["value"] < total[a]["value"]]
            prev = ym(f["year"] - 1, 12)
            dec_to_jan = total[yr[0]]["value"] - total[prev]["value"] if prev in actual else None
            ok = not falls
            out.matches.append({"slide": f["slide"], "figure": "grew every month", "period": f"{yr[0]} to {yr[-1]}",
                                "books": "falls in " + (", ".join(falls) if falls else "none"),
                                "basis": "'Total revenue' month on month", "gap_rel_pct": None, "within_tolerance": ok,
                                "december_to_january_change": None if dec_to_jan is None else round(dec_to_jan, 2)})
            out.claim(f"The deck says revenue grew every month in {f['year']} (slide {f['slide']}); in the P&L 'Total revenue' "
                      + ("rose in every month from " + yr[0] + " to " + yr[-1] if ok else "fell in " + ", ".join(falls))
                      + (f", and changed by {money(dec_to_jan)} from {prev} to {yr[0]}." if dec_to_jan is not None else "."),
                      None, "info" if ok else "flag", CHECK_REVENUE,
                      [src_file(room.deck_uri, f"slide {f['slide']}", f["quote"])]
                      + [src_file(MODEL, total[m]["ref"], quote_of(total[m]["value"])) for m in [prev, yr[0]] if m in actual])

    # Second read of the books: the revenue rows, as a cross-check, never as a claim here.
    cross = None
    if rr:
        rel, sums = rr
        line, lname = (recurring, rec_label) if rel.endswith("customer_mrr_by_month.csv") else (total, "Total revenue")
        diffs = {m: round(sums.get(m, 0) - line[m]["value"], 2) for m in actual}
        cross = {"file": rel, "pnl_line": lname, "months_compared": len(actual),
                 "months_differing_by_more_than_1_dollar": {m: d for m, d in diffs.items() if abs(d) > 1}}
    label = "flag" if any(c["severity"] == "flag" for c in out.claims) else "clear"
    res = finish(out, label, f"{sum(c['severity'] == 'flag' for c in out.claims)} of {len(out.claims)} claims are flags.")
    res["first_revenue_month"] = first_rev
    res["as_of"] = as_of
    res["revenue_rows_cross_check"] = cross
    return res


def describe_months(months: list[str], line: dict) -> dict:
    vals = [line[m]["value"] for m in months]
    peak_i = max(range(len(vals)), key=lambda i: vals[i])
    falls = [months[i] for i in range(1, len(vals)) if vals[i] < vals[i - 1]]
    after = line[_next(months[-1])]["value"]
    if falls or vals[peak_i] > after:
        text = (f"hold a peak of {money(vals[peak_i])} in {months[peak_i]}; revenue then fell "
                f"{100 * (1 - after / vals[peak_i]):.1f}% to {money(after)} in the chart's first month")
    else:
        text = f"hold the climb from {money(vals[0])} to {money(vals[-1])}, with no fall"
    return {"text": text, "peak_month": months[peak_i], "peak": vals[peak_i], "first": vals[0], "last": vals[-1],
            "falls_in": falls, "chart_first_month_value": after,
            "fall_from_peak_pct": round(100 * (1 - after / vals[peak_i]), 1),
            "peak_above_chart_first_pct": round(100 * (vals[peak_i] / after - 1), 1)}


def _next(m: str) -> str:
    y, mo = map(int, m.split("-"))
    return ym(y + (mo == 12), 1 if mo == 12 else mo + 1)


def finish(out: Out, label: str, reason: str) -> dict:
    return {"framework": "deck_vs_books", "label": label, "label_reason": reason,
            "claims": out.claims, "gaps": out.gaps, "matches": out.matches, "unmatched": out.unmatched}


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("room", type=Path)
    ap.add_argument("--deck", type=Path)
    ap.add_argument("--hide", action="append", default=[])
    ap.add_argument("--out", type=Path)
    a = ap.parse_args(argv)
    res = run(Room(a.room, a.deck, set(a.hide)))
    print(f"room: {a.room}   deck: {a.deck or DECK}   hidden: {a.hide or 'none'}")
    print(f"label: {res['label']}   ({res['label_reason']})")
    if res.get("first_revenue_month"):
        print(f"as of: {res['as_of']}   first revenue month in the books: {res['first_revenue_month']}")
    print("\nmatches:")
    for m in res["matches"]:
        tick = "ok  " if m["within_tolerance"] else "MISS"
        gap = "" if m["gap_rel_pct"] is None else f"  gap {m['gap_rel_pct']}%"
        print(f"  {tick} slide {m['slide']:>2}  {m['figure']:<28} {m['period'] or '':<22} books: {m['books']}  [{m['basis']}]{gap}")
    if res["unmatched"]:
        print("\nunmatched:")
        for u in res["unmatched"]:
            print(f"  slide {u['slide']}: {u['kind']} {u.get('figure', '')}: {u['why']}")
    print("\nclaims:")
    for c in res["claims"]:
        print(f"  {c['id']} [{c['severity']}] {c['check']['name']}: {c['text']}")
    if res["gaps"]:
        print("\ngaps:")
        for g in res["gaps"]:
            print(f"  {g['what']} ({g['why']}); would need: {g['would_need']}")
    if res.get("revenue_rows_cross_check"):
        x = res["revenue_rows_cross_check"]
        print(f"\nrevenue rows cross-check: {x['file']} summed by month against '{x['pnl_line']}', "
              f"{x['months_compared']} months, differing by more than $1: {x['months_differing_by_more_than_1_dollar'] or 'none'}")
    if a.out:
        a.out.parent.mkdir(parents=True, exist_ok=True)
        a.out.write_text(json.dumps(res, indent=2, default=str) + "\n")
        print(f"\nwrote {a.out}")


if __name__ == "__main__":
    sys.exit(main())
