#!/usr/bin/env python3
"""Spike: the accounting-mislabels sub-checks (FRAMEWORKS.md entry 1.12).

Throwaway code. It exists to settle behaviour before the real module
(scans/redflags/mislabels.py) is built test first. It reads a data room,
runs each sub-check, prints a table, and writes the result as JSON.

    python3 mislabels_spike.py                    # every room below, writes out/
    python3 mislabels_spike.py --room <room_dir>  # one room
    python3 mislabels_spike.py --schema-dir <dir> # also validate with the skeleton's
                                                  # schema.claim and schema.check_quotes

It writes two test rooms under test_rooms/ first: a copy of the Hollis Voice
room with a resale line (option on) and a plain copy (option off). It never
writes inside data_room/.

Packages: openpyxl (data_room/requirements.txt). jsonschema is optional.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[3]
DATA = REPO / "data_room"
FRAMEWORKS = REPO / "docs/ongoing/FRAMEWORKS.md"
TOL = 0.01  # match_tolerance, FRAMEWORKS.md "Settings"; schema/settings.py DEFAULTS
TODAY = dt.date.today().isoformat()
SCAN = "redflags"

CUST_HOLLIS = "03_metrics/customer_mrr_by_month.csv"
CUST_BRASS = "03_metrics/revenue_by_customer.csv"
METRICS = "03_metrics/monthly_metrics.csv"
LEDGER = "02_financials/expense_ledger.csv"
MODEL = "02_financials/financial_model.xlsx"

PIECE = {
    "mrr_column_vs_subscription_rows": "traction_and_revenue",
    "billings_vs_recognised_revenue": "traction_and_revenue",
    "gross_vs_net_revenue": "traction_and_revenue",
    "contractor_spend_vs_headcount_and_burn": "burn_and_runway",
    "capitalised_spend_vs_pnl_and_cash": "burn_and_runway",
    "deferred_revenue_vs_earned": "traction_and_revenue",
    "one_time_fees_in_recurring_metrics": "traction_and_revenue",
}

CONTRACTOR_RE = re.compile(
    r"\b(contractors?|consultants?|consulting|freelance|freelancers?|agency|staffing|"
    r"contract engineering|statement of work|sow)\b", re.I)
CONTRACTOR_RE_1_12 = re.compile(r"(contractor|consultant|freelance|agency|staffing)", re.I)
ASSET_RE = re.compile(r"(capitali[sz]ed|asset|balance sheet|prepaid)", re.I)
PASS_TYPE_RE = re.compile(r"(resale|pass.?through|reimburs)", re.I)
PASS_COL_RE = re.compile(r"(remitted|pass.?through|paid_to_customer|on_behalf)", re.I)
PASS_LEDGER_RE = re.compile(r"(remittance|resale|pass.?through|at cost)", re.I)


def money(x: float) -> str:
    sign = "-" if x < 0 else ""
    x = abs(x)
    return sign + (f"${x:,.0f}" if abs(x - round(x)) < 0.005 else f"${x:,.2f}")


def pct(x: float) -> str:
    return f"{100 * x:.1f}%"


def cell_text(v) -> str:
    """How schema.readers.text_at renders a cell: str(value)."""
    return "" if v is None else str(v)


def num(s) -> float:
    if s is None or s == "":
        return 0.0
    return float(s)


# ---------------------------------------------------------------------------
# Reading a room. The real module uses schema/readers.py instead.
# ---------------------------------------------------------------------------
class Room:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.name = self.path.name
        self._csv: dict[str, list] = {}
        self._wb = None

    def has(self, rel: str) -> bool:
        return (self.path / rel).exists()

    def rows(self, rel: str) -> list[tuple[int, dict, str]]:
        """(file row number, row dict, row text). Row 1 is the header."""
        if rel not in self._csv:
            lines = (self.path / rel).read_text(encoding="utf-8").splitlines()
            header = next(csv.reader([lines[0]]))
            out = []
            for i, line in enumerate(lines[1:], start=2):
                vals = next(csv.reader([line]))
                out.append((i, dict(zip(header, vals)), line))
            self._csv[rel] = out
        return self._csv[rel]

    def header(self, rel: str) -> list[str]:
        return (self.path / rel).read_text(encoding="utf-8").splitlines()[0].split(",")

    def header_text(self, rel: str) -> str:
        return (self.path / rel).read_text(encoding="utf-8").splitlines()[0]

    @property
    def wb(self):
        if self._wb is None:
            self._wb = openpyxl.load_workbook(self.path / MODEL, data_only=True)
        return self._wb

    def named(self, name: str):
        """(locator 'Sheet!B9', value) for a defined name, or None."""
        if not self.has(MODEL) or name not in self.wb.defined_names:
            return None
        sheet, ref = next(iter(self.wb.defined_names[name].destinations))
        ref = ref.replace("$", "")
        return f"{sheet}!{ref}", self.wb[sheet][ref].value

    def pnl(self):
        """{'months': [(label, status, col letter)], 'rows': {label: (row number, section)}}."""
        ws = self.wb["P&L"]
        hdr = None
        for r in range(1, 12):
            if ws.cell(r, 1).value == "Line":
                hdr = r
                break
        months = []
        for c in range(3, ws.max_column + 1):
            m = ws.cell(hdr, c).value
            if m:
                months.append((str(m), ws.cell(hdr + 1, c).value, ws.cell(hdr, c).column_letter))
        rows = {}
        for r in range(hdr + 2, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if label:
                rows[str(label)] = (r, ws.cell(r, 2).value)
        return {"months": months, "rows": rows, "ws": ws}

    def pnl_value(self, label: str, month: str):
        p = self.pnl()
        if label not in p["rows"]:
            return None
        r = p["rows"][label][0]
        for m, _status, col in p["months"]:
            if m == month:
                return f"P&L!{col}{r}", p["ws"][f"{col}{r}"].value
        return None

    def actual_months(self) -> list[str]:
        return [m for m, s, _ in self.pnl()["months"] if s == "Actual"]

    def first_plan_month(self) -> str | None:
        plans = [m for m, s, _ in self.pnl()["months"] if s == "Plan"]
        return plans[0] if plans else None


def layout(room: Room) -> dict:
    """Column names for the two room shapes the repo has today."""
    if room.has(CUST_BRASS):
        return {"company": "brasswick", "cust": CUST_BRASS, "rev": "subscription_revenue", "fee": None,
                "bill": "billings", "ptype": None, "metrics_mrr": None}
    return {"company": "hollis", "cust": CUST_HOLLIS, "rev": "mrr", "fee": "one_time_fee",
            "bill": "billings", "ptype": "product_type", "metrics_mrr": "mrr"}


# ---------------------------------------------------------------------------
# Output shapes: the shared Claim and Gap (FRAMEWORKS.md, "The shared Claim").
# ---------------------------------------------------------------------------
class Out:
    def __init__(self):
        self.n = 0
        self.claims: list[dict] = []
        self.gaps: list[dict] = []
        self.checks: list[dict] = []
        self.lines: list[dict] = []
        self.links: list[dict] = []

    def claim(self, check, text, value, severity, sources):
        self.n += 1
        cid = f"{SCAN}-mislabels-{self.n:03d}"
        self.claims.append({
            "id": cid, "scan": SCAN, "text": text, "value": value, "severity": severity,
            "check": {"kind": "script", "name": check, "model": None}, "confidence": "computed",
            "sources": sources, "conflicts_with": [], "corroborated_by": [],
        })
        return cid

    def gap(self, check, what, why, would_need):
        self.gaps.append({"what": what, "why": why, "would_need": would_need, "piece": PIECE[check]})

    def row(self, check, label, label_claimed, really_is, file, column, claim_ids):
        self.checks.append({"check_name": check, "defined_in": "1.12", "scope": "after_saturday", "label": label,
                            "label_claimed": label_claimed, "really_is": really_is, "file": file,
                            "column": column, "claim_ids": claim_ids})


def src(uri, locator, quote):
    return {"kind": "file", "uri": uri, "locator": locator, "start_s": None, "end_s": None,
            "quote": quote, "retrieved_at": TODAY}


def csv_src(room: Room, rel: str, rowno: int):
    text = (room.path / rel).read_text(encoding="utf-8").splitlines()[rowno - 1]
    return src(rel, f"row {rowno}", text)


def cell_src(room: Room, loc_value):
    loc, value = loc_value
    return src(MODEL, loc, cell_text(value))


def val(number, unit, label, period):
    return {"number": round(float(number), 2), "unit": unit, "label": label, "period": period}


# ---------------------------------------------------------------------------
# Shared sums
# ---------------------------------------------------------------------------
def by_month_customer(room: Room, lay: dict):
    """month -> list of (rowno, row) for the customer file."""
    out: dict[str, list] = {}
    for rowno, r, _ in room.rows(lay["cust"]):
        out.setdefault(r["month"], []).append((rowno, r))
    return out


def is_sub(r: dict, lay: dict) -> bool:
    return lay["ptype"] is None or r[lay["ptype"]] == "Subscription"


def sub_sums(room: Room, lay: dict):
    rev, bill = {}, {}
    for m, rows in by_month_customer(room, lay).items():
        rev[m] = sum(num(r[lay["rev"]]) for _, r in rows if is_sub(r, lay))
        bill[m] = sum(num(r[lay["bill"]]) for _, r in rows if is_sub(r, lay))
    return rev, bill


def metrics_by_month(room: Room):
    return {r["month"]: (rowno, r) for rowno, r, _ in room.rows(METRICS)}


def ledger_totals(room: Room):
    tot: dict[str, float] = {}
    for _, r, _ in room.rows(LEDGER):
        tot[r["month"]] = tot.get(r["month"], 0.0) + num(r["amount_usd"])
    return tot


def tie_gaps(room: Room):
    tot = ledger_totals(room)
    out = {}
    for m in room.actual_months():
        cor = room.pnl_value("Total cost of revenue", m)[1]
        opx = room.pnl_value("Total operating expenses", m)[1]
        out[m] = tot.get(m, 0.0) - (cor + opx)
    return out, tot


# ---------------------------------------------------------------------------
# a. mrr_column_vs_subscription_rows
# ---------------------------------------------------------------------------
def check_a(room: Room, lay: dict, o: Out):
    name = "mrr_column_vs_subscription_rows"
    if not room.has(METRICS) or "mrr" not in room.header(METRICS):
        o.row(name, "not_run", "MRR", "not_checked", METRICS, "mrr", [])
        return
    met = metrics_by_month(room)
    as_of = max(met)
    m_row, m = met[as_of]
    shown = num(m["mrr"])

    if room.has(lay["cust"]):
        bm = by_month_customer(room, lay)
        sub = {t: sum(num(r["mrr"]) for _, r in rows if r["product_type"] == "Subscription") for t, rows in bm.items()}
        flagged = [t for t in sorted(met) if sub.get(t) and (num(met[t][1]["mrr"]) - sub[t]) / sub[t] > TOL]
        file_used, col = lay["cust"], "product_type, mrr"
        others = [(rn, r) for rn, r in bm.get(as_of, []) if r["product_type"] not in ("Subscription", "Onboarding")
                  and num(r["mrr"]) > 0]
        sub_now = sub[as_of]
        # Group the non-subscription rows by product type; cite the largest row of each type.
        groups: dict[str, list] = {}
        for rn, r in others:
            groups.setdefault(r["product_type"], []).append((rn, r))
        other_srcs, parts = [], []
        for ptype, rows_ in sorted(groups.items()):
            rn, top = max(rows_, key=lambda x: num(x[1]["mrr"]))
            other_srcs.append(csv_src(room, lay["cust"], rn))
            total = sum(num(r["mrr"]) for _, r in rows_)
            parts.append(f"{top['customer_name']} ({ptype}, {money(total)})" if len(rows_) == 1
                         else f"{len(rows_)} {ptype} rows ({money(total)})")
        other_desc = "; ".join(parts)
        other_sum = sum(num(r["mrr"]) for _, r in others)
        # Cite the model's ARR_Build row for the same line too, so this claim shares a source with
        # entry 1.2's pilot-as-ARR flag. That shared source is the link (04_architecture.md, D3).
        if room.has(MODEL) and "ARR_Build" in room.wb.sheetnames:
            ws = room.wb["ARR_Build"]
            for r in range(1, ws.max_row + 1):
                if ws.cell(r, 2).value in groups and num(ws.cell(r, 5).value) > 0:
                    other_srcs += [src(MODEL, f"ARR_Build!A{r}", cell_text(ws.cell(r, 1).value)),
                                   src(MODEL, f"ARR_Build!E{r}", cell_text(ws.cell(r, 5).value))]
    else:
        # Fallback in entry 1.12: the model's ARR_Build Product type column and its monthly history.
        o.gap(name, "Not in the room: a revenue export by product type", "file missing",
              "monthly revenue by customer with a product type column (customer_mrr_by_month.csv)")
        if not room.has(MODEL) or "ARR_Build" not in room.wb.sheetnames:
            o.row(name, "cannot_check", "MRR", "not_checked", METRICS, "mrr", [])
            return
        ws = room.wb["ARR_Build"]
        sub, hist_row = {}, {}
        others, other_srcs = [], []
        in_hist = False
        for r in range(1, ws.max_row + 1):
            a, b, c, d = (ws.cell(r, i).value for i in (1, 2, 3, 4))
            if a == "Month":
                in_hist = True
                continue
            if in_hist and isinstance(a, str) and re.match(r"\d{4}-\d{2}$", a):
                sub[a] = num(c)
                hist_row[a] = r
            elif not in_hist and b not in (None, "Product type", "Subscription", "Onboarding") and num(d) > 0:
                others.append((r, a, b, num(d)))
        flagged = [t for t in sorted(met) if sub.get(t) and (num(met[t][1]["mrr"]) - sub[t]) / sub[t] > TOL]
        file_used, col = MODEL, "ARR_Build Product type, MRR"
        sub_now = sub.get(as_of, 0)
        for r, a, b, d in others:
            other_srcs += [src(MODEL, f"ARR_Build!A{r}", a), src(MODEL, f"ARR_Build!B{r}", b)]
        if as_of in hist_row:
            other_srcs.append(src(MODEL, f"ARR_Build!C{hist_row[as_of]}", cell_text(ws.cell(hist_row[as_of], 3).value)))
        other_desc = "; ".join(f"{a} ({b}, {money(d)})" for _, a, b, d in others)
        other_sum = sum(d for *_, d in others)

    gap = shown - sub_now
    m_src = csv_src(room, METRICS, m_row)
    if not flagged:
        cid = o.claim(name, f"The mrr column in monthly_metrics.csv for {as_of}, {money(shown)}, matches the "
                            f"subscription rows ({money(sub_now)}) within {pct(TOL)}.",
                      val(sub_now, "USD_per_month", "mrr", as_of), "info", [m_src])
        o.row(name, "clear", "MRR", "nothing_found", file_used, col, [cid])
        return
    explained = abs(gap - other_sum) <= TOL * max(gap, 1)
    sub_words = "The subscription rows add to" if room.has(lay["cust"]) else "The model's ARR_Build shows subscription MRR of"
    text = (f"monthly_metrics.csv shows MRR of {money(shown)} for {as_of}. {sub_words} "
            f"{money(sub_now)}. The {money(gap)} difference is {pct(gap / sub_now)} of subscription MRR"
            + (f" and equals the non-subscription lines: {other_desc}." if explained and other_desc else ".")
            + f" Months over tolerance: {', '.join(flagged)}.")
    cid = o.claim(name, text, val(sub_now, "USD_per_month", "mrr", as_of), "flag", [m_src] + other_srcs)
    kind = "pilot" if "Pilot" in other_desc else ("pass_through" if PASS_TYPE_RE.search(other_desc) else "one_time_fee")
    o.row(name, "flag", "MRR", kind, file_used, col, [cid])
    o.lines.append({"locator": f"{METRICS} row {m_row}", "label_shown": "mrr", "kind": "pilot" if kind == "pilot" else "pass_through", "claim_id": cid})
    # Link to entry 1.2's pilot-as-ARR flag: same cause, through a shared source row.
    shared = [s["uri"] + " " + s["locator"] for s in other_srcs]
    o.links.append({"claim_id": cid, "same_cause_as_check": "arr_from_subscription_rows",
                    "shared_sources": shared, "pilot_arr_counted_cell": "ARR_TOTAL"})
    # A P&L row labelled recurring that holds the same non-subscription money.
    if room.has(MODEL):
        for label, (r, _sec) in room.pnl()["rows"].items():
            if "recurring" in label.lower():
                v = room.pnl_value(label, as_of)
                if v and abs(num(v[1]) - (sub_now + other_sum)) <= TOL * sub_now and other_sum > 0:
                    o.lines.append({"locator": v[0], "label_shown": label, "kind": "pilot" if kind == "pilot" else "pass_through", "claim_id": cid})


# ---------------------------------------------------------------------------
# b. billings_vs_recognised_revenue
# ---------------------------------------------------------------------------
def check_b(room: Room, lay: dict, o: Out):
    name = "billings_vs_recognised_revenue"
    if not room.has(lay["cust"]) or lay["bill"] not in room.header(lay["cust"]):
        o.gap(name, "Not in the room: an invoices or billings export", "file missing",
              "monthly invoices or billings by customer")
        o.row(name, "cannot_check", "Revenue", "not_checked", lay["cust"], "billings, mrr", [])
        return
    rev, bill = sub_sums(room, lay)
    months = [m for m in room.actual_months() if m in rev]
    fired, maxdiff = [], 0.0
    for t in months:
        loc, shown = room.pnl_value("Subscription revenue", t)
        shown = num(shown)
        maxdiff = max(maxdiff, abs(shown - rev[t]))
        if abs(shown - bill[t]) < abs(shown - rev[t]) and abs(shown - rev[t]) / rev[t] > TOL:
            fired.append((t, loc, shown))
    arr_fired = None
    arr = room.named("ARR_TOTAL")
    as_of = months[-1]
    if arr:
        a12 = num(arr[1]) / 12
        if abs(a12 - bill[as_of]) < abs(a12 - rev[as_of]) and abs(a12 - rev[as_of]) / rev[as_of] > TOL:
            arr_fired = (arr, a12)
    spikes = [t for t in months if abs(bill[t] - rev[t]) / rev[t] > TOL]
    ids = []
    for t, loc, shown in fired:
        ids.append(o.claim(name, f"The P&L row \"Subscription revenue\" for {t} is {money(shown)}. Subscription "
                                 f"billings in {lay['cust'].split('/')[-1]} add to {money(bill[t])} and subscription "
                                 f"revenue adds to {money(rev[t])}. The P&L figure sits on the billings side.",
                           val(rev[t], "USD", "other", t), "flag", [src(MODEL, loc, cell_text(room.pnl_value('Subscription revenue', t)[1]))]))
    if arr_fired:
        (loc, v), a12 = arr_fired
        ids.append(o.claim(name, f"ARR_TOTAL divided by 12 is {money(a12)}, nearer {as_of} billings ({money(bill[as_of])}) "
                                 f"than {as_of} subscription revenue ({money(rev[as_of])}).",
                           val(12 * rev[as_of], "USD", "arr", as_of), "flag", [src(MODEL, "ARR_TOTAL", cell_text(v))]))
    if ids:
        o.row(name, "flag", "Revenue", "billings", lay["cust"], f"{lay['bill']}, {lay['rev']}", ids)
        return
    loc, shown = room.pnl_value("Subscription revenue", as_of)
    cid = o.claim(name, f"The P&L row \"Subscription revenue\" matches subscription revenue in the customer file in "
                        f"all {len(months)} actual months (largest difference {money(maxdiff)}). Billings differ from "
                        f"revenue by more than {pct(TOL)} in {len(spikes)} of those months, which annual plans paid up "
                        f"front explain on their own.",
                  val(shown, "USD", "other", as_of), "info",
                  [src(MODEL, loc, cell_text(shown)), src(lay["cust"], "row 1", room.header_text(lay["cust"]))])
    o.row(name, "clear", "Revenue", "nothing_found", lay["cust"], f"{lay['bill']}, {lay['rev']}", [cid])


# ---------------------------------------------------------------------------
# c. gross_vs_net_revenue
# ---------------------------------------------------------------------------
def check_c(room: Room, lay: dict, o: Out):
    name = "gross_vs_net_revenue"
    gap_what = "Not in the room: a revenue export that separates pass-through money"
    if not room.has(lay["cust"]):
        o.gap(name, gap_what, "file missing", "a revenue export with a pass-through or resale column, and the ledger lines paid on at cost")
        o.row(name, "cannot_check", "Revenue", "not_checked", lay["cust"], "", [])
        return
    hdr = room.header(lay["cust"])
    bm = by_month_customer(room, lay)
    pt_col = next((h for h in hdr if PASS_COL_RE.search(h)), None)
    fee_col = next((h for h in hdr if re.search(r"fee", h) and h != pt_col and h != lay["fee"]), None)
    pt, incl, excl, desc = {}, {}, {}, ""
    if pt_col and fee_col:
        for t, rows in bm.items():
            pt[t] = sum(num(r[pt_col]) for _, r in rows)
            excl[t] = sum(num(r[fee_col]) for _, r in rows)
            incl[t] = excl[t] + pt[t]
        desc, column = f"column {pt_col}", f"{pt_col}, {fee_col}"
    elif lay["ptype"]:
        types = {r[lay["ptype"]] for _, rows in bm.items() for _, r in rows}
        pt_types = {x for x in types if PASS_TYPE_RE.search(x)}
        if pt_types:
            for t, rows in bm.items():
                pt[t] = sum(num(r["mrr"]) for _, r in rows if r[lay["ptype"]] in pt_types)
                excl[t] = sum(num(r["mrr"]) for _, r in rows if r[lay["ptype"]] not in pt_types)
                incl[t] = excl[t] + pt[t]
            desc, column = f"product type {', '.join(sorted(pt_types))}", "product_type, mrr"
    if not pt or not any(pt.values()):
        types = sorted({r[lay["ptype"]] for _, rows in bm.items() for _, r in rows}) if lay["ptype"] else []
        o.gap(name, gap_what, "not stated in the file",
              "a revenue export with a pass-through or resale column, and the ledger lines paid on at cost")
        o.row(name, "not_run", "Revenue", "not_checked", lay["cust"], ", ".join(types) or "", [])
        return
    # Which revenue-labelled figures hold the pass-through money? Same "nearer" test as kind b.
    months = sorted(t for t in pt if pt[t] > 0)
    as_of = max(bm)
    candidates = []
    if room.has(MODEL):
        for label, (r, sec) in room.pnl()["rows"].items():
            if sec == "Revenue" and not label.lower().startswith("total"):
                candidates.append(("P&L", label, lambda t, l=label: room.pnl_value(l, t)))
    met = metrics_by_month(room) if room.has(METRICS) else {}
    for h in (room.header(METRICS) if room.has(METRICS) else []):
        if re.search(r"(revenue|mrr)", h) and not h.startswith("total"):
            candidates.append(("metrics", h, lambda t, h=h: (f"{METRICS} row {met[t][0]}", met[t][1][h]) if t in met else None))
    holders = []
    for kind, label, get in candidates:
        v = get(as_of)
        if v is None or v[1] in (None, ""):
            continue
        f = num(v[1])
        if abs(f - incl[as_of]) <= TOL * incl[as_of] and abs(f - incl[as_of]) < abs(f - excl[as_of]):
            holders.append((kind, label, v))
    ltm = [t for t in sorted(bm) if t <= as_of][-12:]
    pt_ltm = sum(pt.get(t, 0) for t in ltm)
    if room.has(MODEL) and "Total revenue" in room.pnl()["rows"]:
        shown_ltm = sum(num(room.pnl_value("Total revenue", t)[1]) for t in ltm)
        shown_label = "P&L Total revenue"
    else:
        h = holders[0] if holders else None
        shown_ltm = sum(num(met[t][1][h[1]]) for t in ltm if t in met) if h and h[0] == "metrics" else 0
        shown_label = f"monthly_metrics.csv {h[1]}" if h else "revenue"
    if room.has(MODEL) and "Total revenue" in room.pnl()["rows"] and not any(k == "P&L" for k, *_ in holders):
        # P&L does not hold it; measure against the metrics figure that does.
        h = next((x for x in holders if x[0] == "metrics"), None)
        if h:
            shown_ltm = sum(num(met[t][1][h[1]]) for t in ltm if t in met)
            shown_label = f"monthly_metrics.csv {h[1]}"
    called_gross = any("gross" in label.lower() for _, label, _ in holders)
    share = pt_ltm / shown_ltm if shown_ltm else 0
    # Ledger lines paid on at cost, as a second source.
    led_src, led_ltm = [], 0.0
    if room.has(LEDGER):
        for rowno, r, _ in room.rows(LEDGER):
            if r["month"] in ltm and PASS_LEDGER_RE.search(r["category"] + " " + r["description"]):
                led_ltm += num(r["amount_usd"])
                if r["month"] == as_of and len(led_src) < 1:
                    led_src.append(csv_src(room, LEDGER, rowno))
    if holders and share > TOL and not called_gross:
        srcs = []
        for kind, label, (loc, v) in holders[:2]:
            if kind == "P&L":
                srcs.append(src(MODEL, loc, cell_text(v)))
            else:
                srcs.append(csv_src(room, METRICS, met[as_of][0]))
        pt_row = next((rn for rn, r in bm[as_of] if (num(r.get(pt_col, 0)) if pt_col else (PASS_TYPE_RE.search(r[lay['ptype']]) and num(r['mrr']))) ), None)
        if pt_row:
            srcs.append(csv_src(room, lay["cust"], pt_row))
        srcs += led_src
        names = ", ".join(f"{k} \"{l}\"" for k, l, _ in holders)
        text = (f"{names} for {as_of} equals the gross amount ({money(incl[as_of])}), including "
                f"{money(pt[as_of])} passed on ({desc}). Over the 12 months to {as_of}, {money(pt_ltm)} of pass-through "
                f"money sits inside {money(shown_ltm)} of {shown_label}, {pct(share)}. Net of it: {money(shown_ltm - pt_ltm)}."
                + (f" The ledger shows {money(led_ltm)} paid on in the same months." if led_ltm else ""))
        cid = o.claim(name, text, val(shown_ltm - pt_ltm, "USD", "other", f"{ltm[0]} to {ltm[-1]}"), "flag", srcs)
        o.row(name, "flag", "Revenue", "pass_through", lay["cust"], column, [cid])
        for kind, label, (loc, v) in holders:
            o.lines.append({"locator": loc, "label_shown": label, "kind": "pass_through", "claim_id": cid})
        return
    o.row(name, "clear", "Revenue", "nothing_found", lay["cust"], column, [])


# ---------------------------------------------------------------------------
# e. contractor_spend_vs_headcount_and_burn
# ---------------------------------------------------------------------------
def contractor_lines(room: Room, regex):
    rows = room.rows(LEDGER)
    vendors = {r["vendor"] for _, r, _ in rows if regex.search(r["vendor"] + " " + r["description"] + " " + r["category"])}
    return [(rn, r) for rn, r, _ in rows if r["vendor"] in vendors], vendors


def check_e(room: Room, lay: dict, o: Out):
    name = "contractor_spend_vs_headcount_and_burn"
    if not room.has(LEDGER):
        o.gap(name, "Not in the room: an expense ledger", "file missing", "the expense ledger with vendor, description and amount")
        o.row(name, "cannot_check", "Headcount, burn", "not_checked", LEDGER, "vendor, description, amount_usd", [])
        return
    ties, tot = tie_gaps(room)
    months = sorted(ties)
    as_of = months[-1]
    worst = max(months, key=lambda t: abs(ties[t]))
    lines, vendors = contractor_lines(room, CONTRACTOR_RE)
    lines_112, _ = contractor_lines(room, CONTRACTOR_RE_1_12)
    spend = {}
    for rn, r in lines:
        spend[r["month"]] = spend.get(r["month"], 0.0) + num(r["amount_usd"])
    staff = room.named("STAFF_NOW") or room.named("EMPLOYEES_NOW")
    ids = []
    off_pnl = [t for t in months if abs(ties[t]) / tot[t] > TOL]
    for t in off_pnl:
        ids.append(o.claim(name, f"Ledger lines for {t} add to {money(tot[t])}; P&L cost of revenue plus operating "
                                 f"expenses differ by {money(ties[t])}.", val(ties[t], "USD", "other", t), "flag",
                           [cell_src(room, room.pnl_value("Total operating expenses", t))]))
    plan_flag = None
    now = spend.get(as_of, 0.0)
    plan_m = room.first_plan_month() if room.has(MODEL) else None
    if now > 0 and plan_m:
        # Which P&L row holds the contractor lines? The row whose as-of actual equals their category's ledger sum.
        cats = {r["category"] for _, r in lines if r["month"] == as_of}
        cat_sum = sum(num(r["amount_usd"]) for _, r, _ in room.rows(LEDGER) if r["month"] == as_of and r["category"] in cats)
        row_label = None
        for label in room.pnl()["rows"]:
            v = room.pnl_value(label, as_of)
            if v and isinstance(v[1], (int, float)) and abs(v[1] - cat_sum) < 0.01:
                row_label = label
                break
        if row_label:
            act = num(room.pnl_value(row_label, as_of)[1])
            planv = room.pnl_value(row_label, plan_m)
            covered = num(planv[1]) - (act - now)
            if covered < now * (1 - TOL):
                srcs = [cell_src(room, room.pnl_value(row_label, as_of)), cell_src(room, planv)]
                srcs += [csv_src(room, LEDGER, rn) for rn, r in lines if r["month"] == as_of][:3]
                if staff:
                    srcs.append(src(MODEL, staff[0], cell_text(staff[1])))
                plan_flag = o.claim(
                    name, f"Contractor lines in the ledger add to {money(now)} in {as_of} ({len([1 for _, r in lines if r['month'] == as_of])} "
                          f"lines from {', '.join(sorted({r['vendor'] for _, r in lines if r['month'] == as_of}))}). They sit in the "
                          f"P&L row \"{row_label}\", {money(act)} in {as_of}. The plan puts {money(num(planv[1]))} on that row "
                          f"from {plan_m}, which leaves {money(max(0, now - max(covered, 0)))} a month of the contractors out of the plan's burn."
                          + (f" The staff count is {staff[1]} ({staff[0]})." if staff else ""),
                    val(now, "USD_per_month", "payroll_monthly", as_of), "flag", srcs)
                ids.append(plan_flag)
    if ids:
        o.row(name, "flag", "Headcount, burn", "contractor_spend", LEDGER, "vendor, description, amount_usd", ids)
        o.gap(name, "Not in the room: a statement of whether the staff count includes contractors", "not stated in the file",
              "a contractor list or an org chart that shows them")
        o._missed_by_1_12 = sorted(vendors - {r["vendor"] for _, r in lines_112})
        return
    text = (f"Ledger lines tie to P&L cost of revenue plus operating expenses in all {len(months)} actual months "
            f"(largest difference {money(abs(ties[worst]))}). "
            + (f"Contractor lines found: {money(now)} in {as_of}." if now else "No ledger line names a contractor, consultant, freelancer, agency or statement of work."))
    srcs = [cell_src(room, room.pnl_value("Total operating expenses", as_of))]
    if staff:
        srcs.append(src(MODEL, staff[0], cell_text(staff[1])))
    cid = o.claim(name, text, val(tot[as_of], "USD", "other", as_of), "info" if not now else "watch", srcs)
    o.row(name, "watch" if now else "clear", "Headcount, burn", "contractor_spend" if now else "nothing_found", LEDGER,
          "vendor, description, amount_usd", [cid])
    o.gap(name, "Not in the room: a list of contractors, or an org chart that shows them", "file missing",
          "a contractor list or contractor agreements")


# ---------------------------------------------------------------------------
# f. capitalised_spend_vs_pnl_and_cash
# ---------------------------------------------------------------------------
def check_f(room: Room, lay: dict, o: Out):
    name = "capitalised_spend_vs_pnl_and_cash"
    bridge = room.named("CASH_BRIDGE_DIFF")
    cash = room.named("CASH_START")
    if not room.has(LEDGER):
        o.gap(name, "Not in the room: an expense ledger", "file missing", "the expense ledger with booked_as and category")
        o.row(name, "cannot_check", "Expense, burn", "not_checked", LEDGER, "booked_as, category", [])
        return
    ties, tot = tie_gaps(room)
    months = sorted(ties)
    as_of = months[-1]
    cap = {}
    for rn, r, _ in room.rows(LEDGER):
        if ASSET_RE.search(r["booked_as"] + " " + r["category"]):
            cap[r["month"]] = cap.get(r["month"], 0.0) + num(r["amount_usd"])
    fired = [t for t in months if cap.get(t, 0) / tot[t] > TOL or abs(ties[t]) / tot[t] > TOL]
    bridge_bad = bridge and cash and abs(num(bridge[1])) / num(cash[1]) > TOL
    booked = sorted({r["booked_as"] for _, r, _ in room.rows(LEDGER)})
    srcs = [src(LEDGER, "row 1", room.header_text(LEDGER))]
    if bridge:
        srcs.append(src(MODEL, bridge[0], cell_text(bridge[1])))
    if cash:
        srcs.append(src(MODEL, cash[0], cell_text(cash[1])))
    if fired or bridge_bad:
        cid = o.claim(name, f"Capitalised or off-P&L spend in {', '.join(fired) or 'none'}; cash bridge difference "
                            f"{money(num(bridge[1])) if bridge else 'n/a'}.", None, "flag", srcs)
        o.row(name, "flag", "Expense, burn", "capitalised_spend", LEDGER, "booked_as, category", [cid])
    else:
        cid = o.claim(name, f"Every ledger line is booked as one of {', '.join(booked)}; none names an asset. The cash "
                            f"bridge difference is {money(num(bridge[1])) if bridge else 'not stated'} against cash of "
                            f"{money(num(cash[1])) if cash else 'not stated'}.",
                      val(num(cash[1]), "USD", "cash", as_of) if cash else None, "info", srcs)
        o.row(name, "clear", "Expense, burn", "nothing_found", LEDGER, "booked_as, category", [cid])
    o.gap(name, "Not in the room: a balance sheet", "file missing", "monthly balance sheets")


# ---------------------------------------------------------------------------
# g. deferred_revenue_vs_earned
# ---------------------------------------------------------------------------
def check_g(room: Room, lay: dict, o: Out):
    name = "deferred_revenue_vs_earned"
    if not room.has(lay["cust"]):
        o.gap(name, "Not in the room: a balance sheet with a deferred revenue line", "file missing",
              "monthly balance sheets, or invoices and revenue by customer")
        o.row(name, "cannot_check", "Revenue, ARR", "not_checked", lay["cust"], "billings, mrr, one_time_fee", [])
        return
    bm = by_month_customer(room, lay)
    as_of = max(bm)
    cum: dict[str, float] = {}
    for t in sorted(bm):
        for _, r in bm[t]:
            earned = num(r[lay["rev"]]) + (num(r[lay["fee"]]) if lay["fee"] else 0)
            cum[r["customer_id"]] = cum.get(r["customer_id"], 0.0) + num(r[lay["bill"]]) - earned
    deferred = {c: v for c, v in cum.items() if v > 0.005}
    deferred_total = sum(deferred.values())
    rev, bill = sub_sums(room, lay)
    ids = []
    for t in [m for m in room.actual_months() if m in rev]:
        loc, shown = room.pnl_value("Subscription revenue", t)
        if num(shown) > rev[t] * (1 + TOL):
            excess = num(shown) - rev[t]
            culprits = [(rn, r) for rn, r in bm[t] if num(r[lay["bill"]]) - num(r[lay["rev"]]) > 0.005]
            match = abs(sum(num(r[lay["bill"]]) - num(r[lay["rev"]]) for _, r in culprits) - excess) <= TOL * excess
            names = ", ".join(r["customer_name"] for _, r in culprits)
            srcs = [src(MODEL, loc, cell_text(shown))] + [csv_src(room, lay["cust"], rn) for rn, _ in culprits[:3]]
            ids.append(o.claim(name, f"The P&L row \"Subscription revenue\" for {t} is {money(num(shown))}. Revenue earned "
                                     f"in the customer file is {money(rev[t])}. The {money(excess)} excess "
                                     + (f"equals the unearned part of the invoices billed that month to {names}." if match else "is not explained by one invoice."),
                               val(rev[t], "USD_per_month", "mrr", t), "flag", srcs))
    arr_sub = None
    if room.has(MODEL) and "ARR_Build" in room.wb.sheetnames:
        ws = room.wb["ARR_Build"]
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == "Subscription ARR":
                arr_sub = (f"ARR_Build!E{r}", ws.cell(r, 5).value)
        if arr_sub and num(arr_sub[1]) > 12 * rev[as_of] * (1 + TOL):
            ids.append(o.claim(name, f"ARR_Build Subscription ARR {money(num(arr_sub[1]))} is above 12 times {as_of} "
                                     f"subscription revenue ({money(12 * rev[as_of])}).", val(12 * rev[as_of], "USD", "arr", as_of),
                               "flag", [src(MODEL, arr_sub[0], cell_text(arr_sub[1]))]))
    dnamed = room.named("DEFERRED_REVENUE_USD")
    if dnamed and abs(num(dnamed[1]) - deferred_total) > TOL * max(deferred_total, 1):
        ids.append(o.claim(name, f"Deferred revenue rebuilt from billings less revenue is {money(deferred_total)} across "
                                 f"{len(deferred)} customers at {as_of}. The model's cash bridge carries {money(num(dnamed[1]))}.",
                           val(deferred_total, "USD", "other", as_of),
                           # Info when a P&L flag already fired: the bridge cell is a correct line for the
                           # Tessmore decoy (Brasswick DC1), and eval.score counts any flag citing it.
                           "info" if ids else "flag", [src(MODEL, dnamed[0], cell_text(dnamed[1]))]))
    if ids:
        o.row(name, "flag", "Revenue", "deferred_revenue", lay["cust"], f"{lay['bill']}, {lay['rev']}", ids)
    else:
        srcs = [src(lay["cust"], "row 1", room.header_text(lay["cust"]))]
        if arr_sub:
            srcs.append(src(MODEL, arr_sub[0], cell_text(arr_sub[1])))
        cid = o.claim(name, f"Billed but not yet earned at {as_of}: {money(deferred_total)} across {len(deferred)} customers. "
                            f"The P&L row \"Subscription revenue\" does not exceed revenue earned in any month"
                            + (f", and ARR_Build Subscription ARR ({money(num(arr_sub[1]))}) is 12 times {as_of} subscription revenue." if arr_sub else "."),
                      val(deferred_total, "USD", "other", as_of), "info", srcs)
        o.row(name, "clear", "Revenue", "nothing_found", lay["cust"], f"{lay['bill']}, {lay['rev']}", [cid])
    if not dnamed:
        o.gap(name, "Not in the room: a balance sheet with a deferred revenue line", "file missing", "monthly balance sheets")
    return deferred


# ---------------------------------------------------------------------------
# h. one_time_fees_in_recurring_metrics
# ---------------------------------------------------------------------------
def check_h(room: Room, lay: dict, o: Out):
    name = "one_time_fees_in_recurring_metrics"
    has_fee_row = room.has(MODEL) and any("one-time" in l.lower() or "onboarding" in l.lower() for l in room.pnl()["rows"])
    if not (lay["fee"] and room.has(lay["cust"])) and not has_fee_row:
        o.row(name, "not_run", "MRR, ARR", "not_checked", lay["cust"], "one_time_fee", [])
        return
    met = metrics_by_month(room)
    fees, sub, file_used = {}, {}, None
    if room.has(lay["cust"]) and lay["fee"]:
        for t, rows in by_month_customer(room, lay).items():
            fees[t] = sum(num(r[lay["fee"]]) for _, r in rows)
            sub[t] = sum(num(r["mrr"]) for _, r in rows if r["product_type"] == "Subscription")
        file_used = lay["cust"]
    else:
        label = next(l for l in room.pnl()["rows"] if "one-time" in l.lower())
        for t in room.actual_months():
            fees[t] = num(room.pnl_value(label, t)[1])
        ws = room.wb["ARR_Build"]
        in_hist = False
        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if a == "Month":
                in_hist = True
            elif in_hist and isinstance(a, str) and re.match(r"\d{4}-\d{2}$", a):
                sub[a] = num(ws.cell(r, 3).value)
        file_used = MODEL
    fired = []
    for t in sorted(met):
        g = num(met[t][1]["mrr"]) - sub.get(t, 0)
        if fees.get(t, 0) > 0 and abs(g - fees[t]) / fees[t] <= TOL:
            fired.append(t)
    arr_on_fee = None
    onboard_cell = None
    if room.has(MODEL) and "ARR_Build" in room.wb.sheetnames:
        ws = room.wb["ARR_Build"]
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 2).value == "Onboarding":
                onboard_cell = (f"ARR_Build!E{r}", ws.cell(r, 5).value, f"ARR_Build!A{r}", ws.cell(r, 1).value)
                if num(ws.cell(r, 5).value) > 0:
                    arr_on_fee = onboard_cell
    as_of = max(met)
    if fired or arr_on_fee:
        cid = o.claim(name, f"One-time fees appear in recurring figures: months {', '.join(fired) or 'none'}; "
                            f"ARR counted on the onboarding row {money(num(arr_on_fee[1])) if arr_on_fee else '$0'}.",
                      None, "flag", [csv_src(room, METRICS, met[as_of][0])])
        o.row(name, "flag", "MRR, ARR", "one_time_fee", file_used, "one_time_fee", [cid])
        return
    g = num(met[as_of][1]["mrr"]) - sub.get(as_of, 0)
    srcs = [csv_src(room, METRICS, met[as_of][0])]
    if onboard_cell:
        srcs += [src(MODEL, onboard_cell[2], cell_text(onboard_cell[3])), src(MODEL, onboard_cell[0], cell_text(onboard_cell[1]))]
    cid = o.claim(name, f"One-time fees of {money(fees.get(as_of, 0))} in {as_of} are not in the mrr column: the gap "
                        f"between mrr and subscription rows is {money(g)}, not the fee total"
                        + (f", and the ARR_Build row \"{onboard_cell[3]}\" counts {money(num(onboard_cell[1]))} of ARR." if onboard_cell else "."),
                  val(fees.get(as_of, 0), "USD", "other", as_of), "info", srcs)
    o.row(name, "clear", "MRR, ARR", "nothing_found", file_used, "one_time_fee", [cid])


CHECKS = [("a", check_a), ("b", check_b), ("c", check_c), ("e", check_e), ("f", check_f), ("g", check_g), ("h", check_h)]
ORDER = {"flag": 3, "watch": 2, "clear": 1, "cannot_check": 0, "not_run": 0}


def run(room_dir: Path) -> dict:
    room = Room(room_dir)
    lay = layout(room)
    o = Out()
    for _k, fn in CHECKS:
        fn(room, lay, o)
    worst = max((c["label"] for c in o.checks), key=lambda x: ORDER[x])
    if all(c["label"] in ("cannot_check", "not_run") for c in o.checks):
        worst = "cannot_check"
    fired = [c["check_name"] for c in o.checks if c["label"] == worst]
    result = {"framework": "accounting_mislabels", "label": worst,
              "label_reason": f"Worst sub-check label is {worst}: {', '.join(fired)}.",
              "claims": o.claims, "gaps": o.gaps, "checks": o.checks, "lines": o.lines}
    return {"room": str(room_dir.relative_to(REPO)) if room_dir.is_relative_to(REPO) else str(room_dir),
            "result": result, "cause_links": o.links,
            "contractor_vendors_missed_by_1_12_keywords": getattr(o, "_missed_by_1_12", [])}


# ---------------------------------------------------------------------------
# Test room with a resale line (the generator option, simulated on CSVs only)
# ---------------------------------------------------------------------------
RESALE_PER_LOCATION = 25  # dollars a month of carrier minutes resold at cost; a spike choice, not data
RESALE_FROM = "2026-04"


def write_test_rooms(src_room: Path, out_dir: Path) -> tuple[Path, Path]:
    off = out_dir / "hollis_voice_resale_off"
    on = out_dir / "hollis_voice_resale_on"
    for d in (off, on):
        shutil.copytree(src_room, d, dirs_exist_ok=True)
    # Option on: a "Resale" row per paying subscription clinic from RESALE_FROM, and the same money in the mrr column.
    cust = on / CUST_HOLLIS
    lines = (src_room / CUST_HOLLIS).read_text(encoding="utf-8").splitlines()
    header = lines[0]
    rows = [next(csv.reader([ln])) for ln in lines[1:]]
    extra: dict[str, float] = {}
    out_rows = []
    for r in rows:
        out_rows.append(r)
        month, cid, name, locs, plan, ptype = r[:6]
        if ptype == "Subscription" and month >= RESALE_FROM:
            amt = RESALE_PER_LOCATION * int(locs)
            out_rows.append([month, cid, name, locs, plan, "Resale", str(amt), "0", str(amt)])
            extra[month] = extra.get(month, 0) + amt
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header.split(","))
    w.writerows(out_rows)
    cust.write_text(buf.getvalue(), encoding="utf-8")
    met_lines = (src_room / METRICS).read_text(encoding="utf-8").splitlines()
    out = [met_lines[0]]
    for ln in met_lines[1:]:
        f = ln.split(",")
        if f[0] in extra:
            mrr = int(f[1]) + int(extra[f[0]])
            f[1], f[2] = str(mrr), str(12 * mrr)
        out.append(",".join(f))
    (on / METRICS).write_text("\n".join(out) + "\n", encoding="utf-8")
    return off, on


# ---------------------------------------------------------------------------
# Validation: JSON Schema from FRAMEWORKS.md, the skeleton's models, the quote check.
# ---------------------------------------------------------------------------
def frameworks_schemas():
    text = FRAMEWORKS.read_text(encoding="utf-8")
    blocks = re.findall(r"```json\n(.*?)```", text, re.S)
    defs = json.loads(blocks[0])["$defs"]
    sec = text[text.index("### 1.12 Accounting mislabels"):text.index("### 1.13 Contract terms")]
    shape = json.loads(re.findall(r"```json\n(.*?)```", sec, re.S)[0])
    return defs, shape


def validate(out: dict, room_dir: Path, schema_dir: str | None) -> list[str]:
    problems = []
    try:
        import jsonschema
        defs, shape = frameworks_schemas()
        # The FRAMEWORKS Gap has no 'piece'; the skeleton added it as optional. Allow it here.
        defs = json.loads(json.dumps(defs))
        defs["Gap"]["properties"]["piece"] = {"anyOf": [{"type": "string"}, {"type": "null"}]}
        shape = dict(shape, **{"$defs": defs})
        jsonschema.validate(out["result"], shape)
    except ImportError:
        problems.append("jsonschema not installed; FRAMEWORKS shape not checked")
    except Exception as e:  # noqa: BLE001
        problems.append(f"FRAMEWORKS 1.12 shape: {str(e).splitlines()[0]}")
    # Quote check, the same rule as schema/check_quotes.py (spacing and case ignored, number forms).
    for c in out["result"]["claims"]:
        for s in c["sources"]:
            hay = quote_haystack(room_dir, s)
            if not any(norm(f) in norm(hay) for f in number_forms(s["quote"]) if norm(f)):
                problems.append(f"{c['id']}: quote not found at {s['uri']} {s['locator']}")
    if schema_dir:
        sys.path.insert(0, schema_dir)
        from schema.claim import Claim, Gap, Panel, ScanResult  # noqa: E402
        from schema.check_quotes import check_claim  # noqa: E402
        r = out["result"]
        panels = [Panel(name="mislabel checks", columns=list(r["checks"][0].keys()) if r["checks"] else [],
                        rows=[[json.dumps(v) if isinstance(v, list) else str(v) for v in c.values()] for c in r["checks"]])]
        sr = ScanResult(scan=SCAN, room=room_dir.name, ran_at=TODAY, inputs=[], claims=[Claim(**c) for c in r["claims"]],
                        gaps=[Gap(**g) for g in r["gaps"]], panels=panels)
        for c in sr.claims:
            _, why = check_claim(c, room_dir)
            problems += [f"{c.id}: skeleton check_quotes: {w}" for w in why]
    return problems


def norm(t: str) -> str:
    return re.sub(r"\s+", " ", t).strip().lower()


def number_forms(q: str) -> set[str]:
    forms = {q}
    bare = q.replace(",", "").replace("$", "").strip()
    try:
        f = float(bare)
    except ValueError:
        return forms
    forms |= {bare, str(f)}
    if f.is_integer():
        forms |= {str(int(f)), f"{int(f):,}"}
    return forms


def quote_haystack(room_dir: Path, s: dict) -> str:
    p = room_dir / s["uri"]
    loc = s["locator"]
    if p.suffix == ".csv":
        return p.read_text(encoding="utf-8").splitlines()[int(loc.split()[1]) - 1]
    wb = openpyxl.load_workbook(p, data_only=True)
    if "!" in loc:
        sh, ref = loc.split("!")
        return cell_text(wb[sh.strip("'")][ref].value)
    sh, ref = next(iter(wb.defined_names[loc].destinations))
    return cell_text(wb[sh][ref.replace("$", "")].value)


# ---------------------------------------------------------------------------
def print_result(out: dict):
    r = out["result"]
    print(f"\n=== {out['room']}  entry label: {r['label']}")
    print(f"{'sub-check':42} {'label':13} really_is")
    for c in r["checks"]:
        print(f"{c['check_name']:42} {c['label']:13} {c['really_is']}")
    for c in r["claims"]:
        print(f"  [{c['severity']}] {c['id']} {c['check']['name']}: {c['text']}")
    for g in r["gaps"]:
        print(f"  [gap] {g['what']} ({g['why']})")
    for ln in r["lines"]:
        print(f"  [line] {ln['locator']} '{ln['label_shown']}' is {ln['kind']} -> {ln['claim_id']}")
    for lk in out["cause_links"]:
        print(f"  [link] {lk['claim_id']} same cause as {lk['same_cause_as_check']} via {lk['shared_sources']}")
    if out["contractor_vendors_missed_by_1_12_keywords"]:
        print(f"  [note] contractor vendors entry 1.12's keyword list misses: {out['contractor_vendors_missed_by_1_12_keywords']}")


def decoy_check(out: dict) -> list[str]:
    """Hollis Voice decoys: no flag or watch claim may cite Brookline Row (C0193), a June 2026 billings row, or the onboarding fee."""
    bad = []
    for c in out["result"]["claims"]:
        if c["severity"] not in ("flag", "watch"):
            continue
        for s in c["sources"]:
            q = s["quote"]
            if ",C0193," in q or (q.startswith("2026-06,") and s["uri"].endswith("customer_mrr_by_month.csv")) or ",Onboarding," in q:
                bad.append(f"{c['id']} cites a decoy row: {q}")
        if c["check"]["name"] == "one_time_fees_in_recurring_metrics":
            bad.append(f"{c['id']} flags the onboarding fee")
    return bad


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--room", action="append")
    ap.add_argument("--schema-dir", default=None)
    args = ap.parse_args(argv)
    off, on = write_test_rooms(DATA / "hollis_voice", HERE / "test_rooms")
    rooms = [Path(r).resolve() for r in args.room] if args.room else [
        DATA / "hollis_voice", DATA / "hollis_voice_seed_drop", DATA / "hollis_voice_clean",
        off, on, DATA / "brasswick"]
    (HERE / "out").mkdir(exist_ok=True)
    for rd in rooms:
        out = run(rd)
        print_result(out)
        probs = validate(out, rd, args.schema_dir)
        if "hollis_voice" in rd.name:
            probs += decoy_check(out)
        print(f"  validation: {'ok' if not probs else '; '.join(probs)}")
        (HERE / "out" / f"{rd.name}.json").write_text(json.dumps(out, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
